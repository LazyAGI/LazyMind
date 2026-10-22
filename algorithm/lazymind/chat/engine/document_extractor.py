from __future__ import annotations

from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from typing import Any, Dict, List

from lazyllm.tools.rag.readers import reader_config_inject as _reader_config_inject  # noqa: F401


@dataclass(frozen=True)
class ExtractedDocumentNode:
    """A lightweight reader result shared by chat and file-resource features."""

    text: str
    metadata: Dict[str, Any]


class LocalDocumentReader:
    """Extract document text without loading the optional RAG runtime.

    This reader owns document decoding, not chunking, embedding or retrieval.
    Native Office extraction does not require an Office-to-PDF service.
    """

    def __call__(self, file_path: str) -> List[ExtractedDocumentNode]:
        import lazyllm

        path = Path(file_path)
        suffix = path.suffix.lower()
        route = lazyllm.globals.config['dynamic_ocr_configs'] or {}
        ocr_type = str(route.get('ocr_type') or '').strip().lower()
        if suffix == '.pdf':
            if ocr_type not in ('', 'none'):
                return self._read_ocr(file_path)
            return self._read_pdf(path)
        if suffix == '.docx':
            return self._read_docx(path)
        if suffix == '.pptx':
            # Official MinerU accepts raw presentations. Other PDF OCR readers
            # must not receive Office bytes when desktop has no converter.
            if ocr_type == 'mineru' and self._uses_official_mineru(route):
                return self._read_ocr(file_path)
            return self._read_pptx(path)
        if suffix == '.doc':
            raise ValueError(
                'Legacy .doc files cannot be decoded locally. Convert the file to .docx or PDF first.'
            )
        raise ValueError(f'Unsupported document type: {suffix or "(no extension)"}')

    @staticmethod
    def _uses_official_mineru(route: dict) -> bool:
        from lazyllm.tools.rag.readers.ocrReader.ocr_service import OcrServiceVariant, resolve_ocr_variant

        # An explicit MinerU type without a URL selects the official service,
        # matching DynamicPDFReader._resolve_route.
        return resolve_ocr_variant('mineru', route.get('ocr_url')) == OcrServiceVariant.ONLINE

    @staticmethod
    def _read_ocr(file_path: str):
        return LocalDocumentReader._get_ocr_reader()(file_path)

    @staticmethod
    @lru_cache(maxsize=1)
    def _get_ocr_reader():
        from lazyllm.tools.rag.readers.ocrReader import DynamicPDFReader
        from lazymind.config import config

        return DynamicPDFReader(
            ocr_url=str(config['ocr_server_url'] or ''),
            image_cache_dir=config['ocr_cache_dir'],
            timeout=3600,
        )

    @staticmethod
    def _read_pdf(path: Path) -> List[ExtractedDocumentNode]:
        import fitz

        nodes: List[ExtractedDocumentNode] = []
        with fitz.open(path) as document:
            for page_index, page in enumerate(document):
                nodes.append(ExtractedDocumentNode(
                    text=page.get_text('text', sort=True) or '',
                    metadata={'page': page_index + 1, 'page_label': str(page_index + 1)},
                ))
        return nodes

    @staticmethod
    def _read_docx(path: Path) -> List[ExtractedDocumentNode]:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        document = Document(path)
        parts: List[str] = []
        for child in document.element.body.iterchildren():
            if child.tag.endswith('}p'):
                text = Paragraph(child, document).text.strip()
                if text:
                    parts.append(text)
            elif child.tag.endswith('}tbl'):
                table = Table(child, document)
                parts.extend(
                    '\t'.join(cell.text.strip() for cell in row.cells)
                    for row in table.rows
                    if any(cell.text.strip() for cell in row.cells)
                )
        return [ExtractedDocumentNode(text='\n'.join(parts), metadata={})]

    @staticmethod
    def _read_pptx(path: Path) -> List[ExtractedDocumentNode]:
        from pptx import Presentation

        nodes: List[ExtractedDocumentNode] = []
        presentation = Presentation(path)
        for slide_index, slide in enumerate(presentation.slides):
            parts: List[str] = []
            for shape in slide.shapes:
                if getattr(shape, 'has_text_frame', False):
                    text = str(getattr(shape, 'text', '') or '').strip()
                    if text:
                        parts.append(text)
                if getattr(shape, 'has_table', False):
                    parts.extend(
                        '\t'.join(cell.text.strip() for cell in row.cells)
                        for row in shape.table.rows
                        if any(cell.text.strip() for cell in row.cells)
                    )
            nodes.append(ExtractedDocumentNode(
                text='\n'.join(parts),
                metadata={'page': slide_index + 1, 'page_label': str(slide_index + 1)},
            ))
        return nodes
