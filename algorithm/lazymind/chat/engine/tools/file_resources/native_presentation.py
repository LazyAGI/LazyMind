"""Create editable native slides, optionally using uploaded template layouts."""
from pathlib import Path


def build_native_pptx(slides: list[dict], output: Path, template: Path | None = None) -> dict:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.shapes import PP_PLACEHOLDER

    if not isinstance(slides, list) or not slides:
        raise ValueError('slides must be a non-empty list')
    for slide in slides:
        if (not isinstance(slide, dict) or not isinstance(slide.get('title'), str)
                or not slide['title'].strip() or not isinstance(slide.get('bullets', []), list)
                or any(not isinstance(item, str) for item in slide.get('bullets', []))):
            raise ValueError('Each slide requires a title and a list of plain-text bullets')
    prs = Presentation(str(template)) if template else Presentation()
    if not template:
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    # Keep master/layout parts and theme, but exclude template sample slides.
    # New slides inherit master backgrounds and logos; arbitrary slide-level
    # decorations are not copied or treated as source content.
    for slide_id in list(prs.slides._sldIdLst):
        prs.part.drop_rel(slide_id.rId)
        prs.slides._sldIdLst.remove(slide_id)
    body_types = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT}
    # Prefer a single content placeholder: multi-column layouts leave unused
    # placeholders and can squeeze a plain-text outline into one narrow column.
    layout = min(prs.slide_layouts, key=lambda item: abs(sum(
        ph.placeholder_format.type in body_types for ph in item.placeholders) - 1))
    fallback_count = 0
    for spec in slides:
        slide = prs.slides.add_slide(layout)
        title = slide.shapes.title
        if title is None:
            title = slide.shapes.add_textbox(int(prs.slide_width * .06), int(prs.slide_height * .05),
                                             int(prs.slide_width * .88), int(prs.slide_height * .13))
        title.text = spec['title']
        body = next((ph for ph in slide.placeholders
                     if ph.placeholder_format.type in body_types and ph.has_text_frame), None)
        if body is None:
            fallback_count += 1
            body = slide.shapes.add_textbox(int(prs.slide_width * .06), int(prs.slide_height * .23),
                                            int(prs.slide_width * .88), int(prs.slide_height * .67))
        frame = body.text_frame
        frame.clear()
        frame.word_wrap = True
        for index, text in enumerate(spec.get('bullets', [])):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = text
            if not body.is_placeholder:
                paragraph.font.size = Pt(22)
        if spec.get('notes'):
            slide.notes_slide.notes_text_frame.text = str(spec['notes'])
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))
    return {'slide_count': len(slides), 'fallback_body_count': fallback_count,
            'template_mode': 'layout' if template else 'none'}
