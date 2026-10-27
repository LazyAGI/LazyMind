import importlib.util
from pathlib import Path
import tempfile
import unittest

from pptx import Presentation

spec = importlib.util.spec_from_file_location('native_presentation', (Path(__file__).parents[2] / 'algorithm')
                                              / 'lazymind/chat/engine/tools/file_resources/native_presentation.py')
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)


class NativePresentationTests(unittest.TestCase):
    def test_template_content_is_replaced_without_overwriting_source(self):
        with tempfile.TemporaryDirectory() as temp:
            source, output = Path(temp) / 'template.pptx', Path(temp) / 'output.pptx'
            template = Presentation()
            template.slides.add_slide(template.slide_layouts[1]).shapes.title.text = 'Sample content'
            template.save(source)
            before = source.read_bytes()
            result = module.build_native_pptx([
                {'title': 'Research', 'bullets': ['Evidence'], 'notes': 'Speaker notes'},
                {'title': 'Conclusion', 'bullets': ['Result']},
            ], output, source)
            actual = Presentation(output)
            self.assertEqual(result['slide_count'], 2)
            self.assertEqual(len(actual.slides), 2)
            self.assertEqual(actual.slides[0].shapes.title.text, 'Research')
            self.assertIn('Evidence', '\n'.join(s.text for s in actual.slides[0].shapes if s.has_text_frame))
            self.assertEqual(actual.slides[0].notes_slide.notes_text_frame.text, 'Speaker notes')
            self.assertEqual(source.read_bytes(), before)
            self.assertEqual(actual.slide_width, template.slide_width)

    def test_plain_presentation_is_editable(self):
        with tempfile.TemporaryDirectory() as temp:
            output = Path(temp) / 'plain.pptx'
            module.build_native_pptx([{'title': 'Plain', 'bullets': ['Editable text']}], output)
            actual = Presentation(output)
            self.assertEqual(actual.slides[0].shapes.title.text, 'Plain')


if __name__ == '__main__':
    unittest.main()
