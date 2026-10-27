import importlib.util
import json
from pathlib import Path

import pytest
import yaml

ROOT = Path(__file__).resolve().parents[4]
spec = importlib.util.spec_from_file_location('github_contributions', ROOT /
    'algorithm/lazymind/chat/engine/tools/file_resources/github_contributions.py')
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)


def scope(**kwargs):
    return module.normalize_scope('https://github.com/acme/project', 'alice',
                                  '2026-01-01', '2026-10-01', **kwargs)


def pr(number=1, author='alice', body='Implemented a feature', **kwargs):
    return {'number': number, 'user': {'login': author}, 'body': body,
            'created_at': '2026-05-01T12:00:00Z', 'state': 'open',
            'title': 'Title is not contribution evidence', 'diff': 'DO NOT READ',
            'head': {'ref': 'feature'}, 'base': {'ref': 'main'}, **kwargs}


def test_filters_author_dates_branch_and_exposes_only_description(tmp_path):
    calls = []
    def api(method, path, **kwargs):
        calls.append((method, path, kwargs))
        return [pr(), pr(2, 'bob'), pr(3, created_at='2025-01-01T00:00:00Z'),
                pr(4, head={'ref': 'other'}), pr(5, body=''), pr()]
    result = module.collect_page(tmp_path, scope(branch='feature'), api=api)
    assert [r['number'] for r in result['records']] == [1, 5]
    assert result['records'][1]['empty_description']
    assert calls[0][0:2] == ('GET', '/repos/acme/project/pulls')
    assert calls[0][2]['allow_redirects'] is False
    record = result['records'][0]
    body = module.read_description(tmp_path, record['source_id'])
    assert body['description'] == 'Implemented a feature'
    assert 'diff' not in json.dumps(result)
    assert 'Title is not contribution evidence' not in ''.join(p.read_text(encoding='utf-8') for p in tmp_path.iterdir())


def test_empty_filtered_page_still_paginates_and_base_filter_is_sent(tmp_path):
    captured = []
    def api(*args, **kwargs):
        captured.append(kwargs['params'])
        return [pr(n, 'bob') for n in range(1, 101)]
    result = module.collect_page(tmp_path, scope(branch='main', branch_kind='base'), api=api)
    assert result['records'] == []
    assert result['next_page'] == 2
    assert captured[0]['base'] == 'main'
    assert captured[0]['per_page'] == 100


def test_description_is_read_to_eof_without_network_and_cache_is_stable(tmp_path):
    text = '这是贡献描述。' * 3000
    result = module.collect_page(tmp_path, scope(), api=lambda *a, **kw: [pr(body=text)])
    record = result['records'][0]
    chunks, offset = [], 0
    while True:
        chunk = module.read_description(tmp_path, record['source_id'], offset)
        chunks.append(chunk['description'])
        if chunk['next_offset'] is None:
            break
        offset = chunk['next_offset']
    assert ''.join(chunks) == text
    repeated = module.collect_page(tmp_path, scope(), api=lambda *a, **kw: [pr(body=text)])
    assert repeated['records'][0]['source_id'] == record['source_id']
    assert len(list(tmp_path.glob(f"{record['source_id']}.json"))) == 1


def test_merged_and_closed_are_distinguished_and_dates_inclusive(tmp_path):
    rows = [pr(1, created_at='2026-01-01T00:00:00Z', merged_at='2026-02-01T00:00:00Z'),
            pr(2, created_at='2026-10-01T23:59:59Z', state='closed'),
            pr(3, created_at='2026-10-02T00:00:00Z')]
    records = module.collect_page(tmp_path, scope(), api=lambda *a, **kw: rows)['records']
    assert [r['state'] for r in records] == ['merged', 'closed']


@pytest.mark.parametrize('repository,author,start,end', [
    ('https://evil.example/acme/project', 'alice', '2026-01-01', '2026-02-01'),
    ('acme/project/tree/main', 'alice', '2026-01-01', '2026-02-01'),
    ('acme/project', 'me', '2026-01-01', '2026-02-01'),
    ('acme/project', '', '2026-01-01', '2026-02-01'),
    ('acme/project', 'alice', '2026-02-01', '2026-01-01'),
    ('acme/project', 'alice', '2026-02-30', '2026-03-01'),
])
def test_invalid_scope_is_rejected(repository, author, start, end):
    with pytest.raises(ValueError):
        module.normalize_scope(repository, author, start, end)


def test_errors_do_not_leak_provider_data(tmp_path):
    def api(*args, **kwargs):
        raise RuntimeError('Authorization: secret-token')
    with pytest.raises(ValueError, match='GITHUB_READ_FAILED') as caught:
        module.collect_page(tmp_path, scope(), api=api)
    assert 'secret-token' not in str(caught.value)
    with pytest.raises(ValueError):
        module.read_description(tmp_path, '../../credentials')


def test_cache_integrity(tmp_path):
    record = module.collect_page(tmp_path, scope(), api=lambda *a, **kw: [pr()])['records'][0]
    (tmp_path / f"{record['source_id']}.json").write_text('{}')
    with pytest.raises(ValueError, match='integrity'):
        module.read_description(tmp_path, record['source_id'])


def test_connected_api_uses_current_session_credentials(monkeypatch):
    import lazyllm
    from lazyllm.common.globals import new_session
    from lazyllm.tools.fs import GitHubRepoFS
    from lazyllm.tools.tool_config_inject import inject_tool_config
    with new_session('ppt-github-test'):
        lazyllm.globals.config['dynamic_fs_auth'] = {}
        with pytest.raises(ValueError, match='connection unavailable'):
            module.connected_api()
        inject_tool_config({'github': 'test-only-token'})
        calls = []
        def fake_api(self, method, path, **kwargs):
            assert self._resolve_dynamic_token() == 'test-only-token'
            calls.append((method, path))
            return []
        monkeypatch.setattr(GitHubRepoFS, '_api', fake_api)
        assert module.connected_api()('GET', '/repos/acme/project/pulls') == []
        assert calls == [('GET', '/repos/acme/project/pulls')]


def test_workflow_declares_cloud_credentials_and_pr_tools_only_in_collection():
    state = yaml.safe_load((ROOT / 'workflows/ppt-workflow/scenario/state.yml').read_text(encoding='utf-8'))
    workflow = yaml.safe_load((ROOT / 'workflows/ppt-workflow/workflow.yaml').read_text(encoding='utf-8'))
    tools = state['steps']['collect_materials']['tools']
    assert {'cloud_files', 'kb', 'ppt_source_connections', 'ppt_collect_github_prs',
            'ppt_read_pr_description'} <= set(tools)
    assert 'cloud_files' not in state['steps']['generate_ppt']['tools']
    declared = set(workflow['tool_scripts'][0]['functions'])
    assert {'ppt_source_connections', 'ppt_collect_github_prs', 'ppt_read_pr_description'} <= declared


def test_workflow_tools_connect_pr_description_to_native_template_publication(monkeypatch, tmp_path):
    from pptx import Presentation
    from lazymind.chat.engine.tools.file_resources import github_contributions as live_module
    tool_spec = importlib.util.spec_from_file_location('ppt_pr_pipeline_test_tools',
        ROOT / 'workflows/ppt-workflow/scripts/tools.py')
    tools = importlib.util.module_from_spec(tool_spec)
    tool_spec.loader.exec_module(tools)
    monkeypatch.setattr(tools, '_conversation_root', lambda: tmp_path)
    calls = []
    def api(method, path, **kwargs):
        calls.append((method, path))
        return [pr(body='完成工作流版本修复，保留历史页面。')]
    monkeypatch.setattr(live_module, 'request_api', lambda _: api)
    found = tools.ppt_collect_github_prs('acme/project', 'alice', '2026-01-01', '2026-10-01')
    evidence = tools.ppt_read_pr_description(found['records'][0]['source_id'])
    source = tmp_path / 'template.pptx'
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = '工作总结'
    prs.save(source)
    original = source.read_bytes()
    monkeypatch.setattr(tools, '_resolve_template_file', lambda _: (source, source.name))
    page = tools.ppt_inspect_template_page(source.name)
    edits = [{'id': slot['id'], 'expected_text': slot['text'],
              'action': 'replace' if slot['empty_paragraph'] else 'keep',
              'text': evidence['description']} for slot in page['slots']]
    published = []
    monkeypatch.setattr(tools, '_save_artifact', lambda **kwargs: published.append(kwargs) or {'ok': True})
    result = tools.ppt_fill_template(source.name, page['sha256'], edits)
    assert result['representation'] == 'native_pptx'
    assert published[0]['key'] == 'native_pptx'
    actual = Presentation(published[0]['value'])
    assert evidence['description'] in '\n'.join(s.text for s in actual.slides[0].shapes if s.has_text_frame)
    assert source.read_bytes() == original
    assert calls == [('GET', '/repos/acme/project/pulls')]


def test_public_closed_author_url_requires_no_credentials(tmp_path):
    link = 'https://github.com/LazyAGI/LazyMind/pulls?q=is%3Apr+state%3Aclosed+author%3ACarlosShaoting'
    selected = module.normalize_scope(link, access_mode='public')
    assert selected['repository'] == 'LazyAGI/LazyMind'
    assert selected['author'] == 'CarlosShaoting'
    assert selected['state'] == 'closed'
    assert selected['start_date'] == selected['end_date'] == ''
    calls = []
    def api(method, path, **kwargs):
        calls.append(kwargs['params'])
        return [pr(1, 'CarlosShaoting', state='closed', merged_at='2026-06-01T00:00:00Z'),
                pr(2, 'CarlosShaoting', state='closed'), pr(3, 'CarlosShaoting')]
    result = module.collect_page(tmp_path, selected, api=api)
    assert calls[0]['state'] == 'closed'
    assert [r['state'] for r in result['records']] == ['merged', 'closed']


def test_public_request_never_uses_ambient_credentials(monkeypatch):
    import requests
    captured = {}
    class Session:
        trust_env = True
        def __enter__(self):
            return self
        def __exit__(self, *args):
            pass
        def get(self, url, **kwargs):
            assert self.trust_env is False
            captured.update(kwargs)
            assert url == 'https://api.github.com/repos/acme/project/pulls'
            return type('Response', (), {'status_code': 200, 'json': lambda self: []})()
    monkeypatch.setattr(requests, 'Session', Session)
    assert module.public_api('GET', '/repos/acme/project/pulls', params={'state': 'closed'}) == []
    assert 'Authorization' not in captured['headers']
    assert captured['allow_redirects'] is False


def test_unsupported_or_conflicting_url_filters_are_not_silently_dropped():
    with pytest.raises(ValueError, match='Unsupported'):
        module.normalize_scope('https://github.com/acme/project/pulls?q=author:alice+label:bug')
    with pytest.raises(ValueError, match='conflicts'):
        module.normalize_scope('https://github.com/acme/project/pulls?q=author:alice', author='bob')
