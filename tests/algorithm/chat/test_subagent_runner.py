"""Tests for lazymind.chat.engine.subagent.runner.

Uses a FakeDB (in-memory) and a FakeAgent (injects known tag sequences) to
drive run_subagent_stream without any real database, LLM, or network calls.
"""
from __future__ import annotations

import asyncio
import base64
import json
import threading
import time
from pathlib import Path
from typing import Any, Dict, List, Optional
from unittest.mock import MagicMock, patch

import pytest

import lazymind.chat.engine.subagent.runner as runner_mod


def test_workflow_script_tool_is_loaded_from_pinned_revision():
    source = 'def create_list_fixtures():\n    return ["one", "two"]\n'
    response = MagicMock()
    response.result = {
        'revision_id': 'revision-1',
        'tree_hash': 'tree-1',
        'files': {
            'scripts/tools.py': base64.b64encode(source.encode()).decode(),
        },
    }
    client = MagicMock()
    client.get_workflow.return_value = response

    with patch('lazymind.workflow_sdk.WorkflowClient', return_value=client):
        tools = runner_mod._resolve_runtime_tools(
            ['create_list_fixtures'],
            {
                'workflow_id': 'test-workflow',
                'revision_id': 'revision-1',
                'tree_hash': 'tree-1',
                'user_id': 'user-1',
            },
        )

    assert [tool.__name__ for tool in tools] == ['create_list_fixtures']
    assert tools[0]() == ['one', 'two']
    client.get_workflow.assert_called_once_with('test-workflow', 'revision-1')


def test_terminal_tools_only_filters_model_tools_without_mutating_runtime_tools():
    def cloud_files():
        pass

    def writer_prepare_workspace():
        pass

    runtime_tools = [cloud_files, writer_prepare_workspace]

    assert runner_mod._model_visible_runtime_tools(runtime_tools, {}) is runtime_tools
    visible = runner_mod._model_visible_runtime_tools(runtime_tools, {
        'terminal_tools_only': True,
        'terminal_tools': ['writer_prepare_workspace'],
    })

    assert runtime_tools == [cloud_files, writer_prepare_workspace]
    assert visible == [writer_prepare_workspace]


def test_large_chinese_tool_result_below_token_limit_stays_inline(tmp_path):
    """A byte-sized Chinese result must not be offloaded before 32K estimated tokens."""
    from lazymind.chat.engine.subagent.context import SubAgentContext

    ctx = SubAgentContext(
        task_id='task-large-zh',
        conversation_id='conv-1',
        agent_type='workflow_step',
        objective='test large result',
        params={},
        workspace_path=str(tmp_path),
        input_slots=[],
        output_slots=[],
        db=None,
        emit=lambda _event: None,
    )
    result = '中' * 22_000  # 66 KB in UTF-8, but only 24,200 estimated tokens.

    assert runner_mod._truncate_tool_result(ctx, result, 'read_file') == result
    assert not (tmp_path / 'large').exists()


def test_tool_result_at_token_limit_uses_shared_spill_reference(tmp_path):
    """A persisted large tool result must use the same store as history compaction."""
    from lazymind.chat.engine.subagent.context import SubAgentContext

    ctx = SubAgentContext(
        task_id='task-large-en',
        conversation_id='conv-1',
        agent_type='workflow_step',
        objective='test large result',
        params={},
        workspace_path=str(tmp_path),
        input_slots=[],
        output_slots=[],
        db=None,
        emit=lambda _event: None,
    )
    result = 'a' * 131_072  # 128 KB and exactly 32,768 estimated tokens.

    rendered = runner_mod._truncate_tool_result(ctx, result, 'read_file')

    assert 'workspace://tool_spills/' in rendered
    assert list((tmp_path / 'tool_spills').glob('read_file_*.txt'))
    assert not (tmp_path / 'large').exists()


def test_durable_and_online_tool_result_spills_share_one_reference(tmp_path):
    from lazymind.chat.engine.agent_runtime.workflow_compactor import (
        make_workflow_history_compactor,
    )
    from lazymind.chat.engine.subagent.context import SubAgentContext

    ctx = SubAgentContext(
        task_id='task-shared-spill',
        conversation_id='conv-1',
        agent_type='workflow_step',
        objective='test shared result spill',
        params={},
        workspace_path=str(tmp_path),
        input_slots=[],
        output_slots=[],
        db=None,
        emit=lambda _event: None,
    )
    result = 'a' * 131_072
    durable_notice = runner_mod._truncate_tool_result(ctx, result, 'read_file')
    compactor = make_workflow_history_compactor(
        max_input_tokens='32K', workspace=str(tmp_path), keep_recent=0,
    )
    prior, _ = compactor([
        {'role': 'assistant', 'content': '', 'tool_calls': [{
            'id': 'call-1', 'function': {'name': 'read_file', 'arguments': '{}'},
        }]},
        {'role': 'tool', 'tool_call_id': 'call-1', 'name': 'read_file', 'content': result},
    ], prefix={'system_prompt': 'workflow system'}, current_input='continue')

    assert prior[1]['content'] == durable_notice
    assert len(list((tmp_path / 'tool_spills').glob('read_file_*.txt'))) == 1


# ---------------------------------------------------------------------------
# In-memory FakeDB
# ---------------------------------------------------------------------------

class FakeDB:
    def __init__(self, task: Optional[Dict[str, Any]] = None):
        self._task = task
        self.steps: List[Dict[str, Any]] = []

    def load_task(self, task_id: str) -> Optional[Dict[str, Any]]:
        return self._task

    def append_step(self, task_id: str, seq: int, role: str, content: Dict[str, Any]) -> None:
        self.steps.append({'task_id': task_id, 'seq': seq, 'role': role, 'content': content})

    def load_steps(self, task_id: str) -> List[Dict[str, Any]]:
        return [s for s in self.steps if s['task_id'] == task_id]

    def max_step_seq(self, task_id: str) -> int:
        relevant = [s['seq'] for s in self.steps if s['task_id'] == task_id]
        return max(relevant) if relevant else -1

    def next_artifact_seq(self, task_id: str, key: str) -> int:
        return 1

    def save_artifact(self, task_id: str, key: str, content_type: str,
                      value: Dict[str, Any], seq: int) -> None:
        pass

    def load_artifacts(self, task_id: str, keys=None) -> List[Dict[str, Any]]:
        return []

    def saved_artifact_keys(self, task_id: str) -> List[str]:
        return []

    def dispose(self) -> None:
        pass


# ---------------------------------------------------------------------------
# Default task fixture
# ---------------------------------------------------------------------------

_DEFAULT_TASK_ID = 'task-001'
_DEFAULT_TASK = {
    'id': _DEFAULT_TASK_ID,
    'conversation_id': 'conv-1',
    'agent_type': 'test',
    'objective': 'do something',
    'params': {'required_output_artifact_keys': ['result']},
    'workspace_path': '/tmp/ws',
    'input_artifact_keys': [],
    'output_artifact_keys': ['result'],
    'mode': 'auto',
}


@pytest.fixture(autouse=True)
def isolated_default_workspace(monkeypatch, tmp_path):
    monkeypatch.setitem(_DEFAULT_TASK, 'workspace_path', str(tmp_path))


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _sse_to_events(raw: str) -> List[Dict[str, Any]]:
    """Parse SSE lines into a list of dicts (skips [DONE] and empty lines)."""
    events = []
    for line in raw.splitlines():
        line = line.strip()
        if line.startswith('data: ') and '[DONE]' not in line:
            events.append(json.loads(line[len('data: '):]))
    return events


async def _collect(gen, on_chunk=None) -> str:
    parts = []
    async for chunk in gen:
        if on_chunk is not None:
            on_chunk(chunk)
        parts.append(chunk)
    return ''.join(parts)


def _install_fake_db(monkeypatch, task=None):
    db = FakeDB(task or {**_DEFAULT_TASK})
    monkeypatch.setattr(
        runner_mod,
        'MemorySubAgentStore',
        lambda task_spec, initial_steps=None, artifacts=None: db,
    )
    return db


def _install_fake_lazyllm(monkeypatch):
    """Patch lazyllm globals/locals init and AutoModel."""
    fake_llm_mod = MagicMock()
    fake_llm_mod.globals._init_sid = lambda sid: None
    fake_llm_mod.locals._init_sid = lambda sid: None
    monkeypatch.setattr(runner_mod, 'lazyllm', fake_llm_mod)
    model = MagicMock()
    model.share.return_value.return_value = json.dumps({
        'completed': True, 'requires_artifact': False, 'artifact_keys': [],
        'reason': 'The requested text result was delivered.',
    })
    monkeypatch.setattr(runner_mod, 'AutoModel', lambda **_: model)
    monkeypatch.setattr(runner_mod, 'inject_model_config', lambda cfg: None)
    monkeypatch.setattr(runner_mod, 'set_context', lambda ctx: None)


def _install_fake_drive(monkeypatch, events, final_value='task done'):
    """Replace AgentExecutor with a deterministic event stream."""
    class FakeExecutor:
        async def stream(self, llm, plan):
            for ev in events:
                yield 'event', ev
            yield 'final', final_value

    monkeypatch.setattr(runner_mod, 'AgentExecutor', FakeExecutor)


def _install_fake_build(monkeypatch):
    """Agent creation is covered by AgentExecutor tests; runner tests replace the executor."""


def _install_fake_translator(monkeypatch):
    """AgentEventFrameTranslator that turns text events into {text:...} frames."""
    class FakeTranslator:
        def __init__(self, query=''):
            self.citation_state: Dict[str, Any] = {}

        def feed(self, item):
            tag = item.get('tag', '')
            if tag == 'text':
                return [{'text': item.get('delta', ''), 'think': None}]
            if tag == 'think':
                return [{'text': None, 'think': item.get('delta', '')}]
            return []

        def finish(self, result):
            return []

    monkeypatch.setattr(runner_mod, 'AgentEventFrameTranslator', FakeTranslator)


def test_created_artifact_task_without_output_slots_cannot_succeed(monkeypatch, tmp_path):
    """Synthetic creation -> runner -> terminal regression; no real model or service."""
    import lazymind.chat.engine.tools.subagent_chat_tools as chat_tools

    monkeypatch.setattr(chat_tools, '_agentic_config', lambda: {'mode': 'manual'})
    created = []
    monkeypatch.setattr(chat_tools, '_write_agent_data', lambda tag, **kw: created.append(kw))
    chat_tools.create_subagent(
        agent_type='document_generation', title='Export document',
        objective='Generate a PDF document and deliver the signed file.',
    )
    task = {**created[0], 'id': created[0]['task_id'], 'workspace_path': str(tmp_path)}
    assert task['output_slots'] == []
    _install_fake_lazyllm(monkeypatch)
    _install_fake_translator(monkeypatch)
    monkeypatch.setattr(
        runner_mod, '_generate_display_plan',
        lambda *_args, **_kwargs: ['Prepare task', 'Execute task', 'Deliver result'],
    )
    model = MagicMock()
    model.share.return_value.return_value = json.dumps({
        'completed': False, 'requires_artifact': True, 'artifact_keys': [],
        'reason': 'The required input and PDF conversion capability were unavailable.',
    })
    monkeypatch.setattr(runner_mod, 'AutoModel', lambda **_: model)
    _install_fake_drive(monkeypatch, [], final_value=(
        'I cannot complete this task: the input file, PDF conversion, and signing '
        'capabilities are unavailable. No artifact was produced.'
    ))

    raw = asyncio.run(_collect(runner_mod.run_subagent_stream(task['id'], task_spec=task)))
    terminal = [event for event in _sse_to_events(raw) if event['type'] in {'done', 'error'}]
    assert len(terminal) == 1
    assert terminal[0]['status'] == 'failed', terminal
    assert terminal[0]['current_phase'] == 'missing_required_artifacts'
    _assert_silent_stream_call(model)
    assert raw.endswith('data: [DONE]\n\n')


@pytest.fixture
def contract_delivery_harness(monkeypatch, tmp_path):
    """Connect the real creation, runner, and parent query paths in memory."""
    import lazymind.chat.engine.tools.subagent_chat_tools as chat_tools

    _install_fake_lazyllm(monkeypatch)
    _install_fake_translator(monkeypatch)
    monkeypatch.setattr(chat_tools, '_agentic_config', lambda: {
        'mode': 'auto', 'conversation_id': 'contract-regression',
    })
    contexts = []
    monkeypatch.setattr(runner_mod, 'set_context', contexts.append)
    monkeypatch.setattr(runner_mod.subagent_tools, 'require_context', lambda: contexts[-1])
    model = MagicMock()
    model.share.return_value.return_value = json.dumps({
        'completed': False, 'requires_artifact': True, 'artifact_keys': [],
        'reason': 'The required contract input and PDF tools were unavailable.',
    })
    monkeypatch.setattr(runner_mod, 'AutoModel', lambda **_: model)
    projected = {}
    observed = {}

    class InMemoryCoreView:
        def get_task_status(self, task_id):
            return projected[task_id]

        def list_tasks_by_conversation(self, _conversation_id):
            return list(projected.values())

    monkeypatch.setattr(chat_tools, 'TaskQueryDB', InMemoryCoreView)

    def run(*, source_path=None, output_slots=None, params=None):
        workspace = tmp_path / ('complete' if source_path else 'blocked')

        class ControlledExecutor:
            async def stream(self, _llm, _plan):
                if source_path:
                    from docx import Document

                    source_text = source_path.read_text(encoding='utf-8')
                    document = Document()
                    document.add_heading('Contract summary', 0)
                    document.add_paragraph(source_text)
                    output_path = workspace / 'contract-summary.docx'
                    document.save(output_path)
                    saved = runner_mod.subagent_tools.save_artifacts([{
                        'key': 'document', 'value': str(output_path),
                        'content_type': 'file',
                    }])
                    assert saved['saved_count'] == 1
                    yield 'final', 'The contract summary document was delivered.'
                else:
                    yield 'final', (
                        'I cannot complete the PDF delivery: the contract input, '
                        'PDF conversion, and signing tools are unavailable. No file was produced.'
                    )

        monkeypatch.setattr(runner_mod, 'AgentExecutor', ControlledExecutor)

        def receive(tag, **created):
            assert tag == 'task_created'
            task_id = created['task_id']
            task = {**created, 'id': task_id, 'workspace_path': str(workspace)}
            observed['task'] = task
            raw = asyncio.run(_collect(runner_mod.run_subagent_stream(task_id, task_spec=task)))
            assert raw.endswith('data: [DONE]\n\n')
            events = _sse_to_events(raw)
            observed['events'] = events
            row = {'task_id': task_id, 'title': created['title'],
                   'status': 'running', 'artifacts': []}
            for event in events:
                if event['type'] == 'artifact':
                    row['artifacts'].append({
                        key: event[key] for key in ('slot', 'content_type', 'value', 'seq')
                    })
                elif event['type'] in {'done', 'error'}:
                    row.update({key: event[key] for key in (
                        'status', 'summary', 'current_phase',
                    ) if key in event})
            projected[task_id] = row

        monkeypatch.setattr(chat_tools, '_write_agent_data', receive)
        result = chat_tools.create_subagent(
            agent_type='document_generation', title='Contract document',
            objective=(
                'Generate and deliver the contract PDF, then sign it.'
                if source_path is None else
                'Read the supplied contract terms and deliver a contract summary document.'
            ),
            params={**(params or {}), **({'source_path': str(source_path)} if source_path else {})},
            output_slots=output_slots,
        )
        return result, observed, chat_tools

    return run


def test_contract_badcase_fails_through_parent_result(contract_delivery_harness):
    parent_result, observed, chat_tools = contract_delivery_harness()
    task = observed['task']
    terminal = _terminal(observed['events'])
    assert task['output_slots'] == []
    assert terminal['status'] == 'failed'
    assert terminal['current_phase'] == 'missing_required_artifacts'
    assert not any(event['type'] == 'artifact' for event in observed['events'])
    assert not list(Path(task['workspace_path']).rglob('*.pdf'))
    assert parent_result['status'] == 'failed'
    assert parent_result['task_status'] == 'failed'
    assert parent_result['failure']['code'] == terminal['current_phase']
    assert parent_result['failure']['message']
    assert chat_tools.get_subagent_status(task['title'])['task']['status'] == 'failed'
    assert chat_tools.get_subagent_artifacts(task['title'])['artifacts'] == []
    print('BADCASE_EVIDENCE=' + json.dumps({
        'case': 'A', 'objective': task['objective'], 'output_slots': task['output_slots'],
        'contract': 'uncontracted artifact goal; semantic completion evaluation',
        'terminal': terminal['status'], 'failure_code': parent_result['failure']['code'],
        'artifact_count': 0, 'parent_status': parent_result['status'],
        'failure_message': parent_result['failure']['message'],
        'agent_final': 'explicit inability; no file produced',
    }, ensure_ascii=False))


def test_contract_document_is_delivered_to_parent(contract_delivery_harness, tmp_path):
    from docx import Document

    source = tmp_path / 'contract-terms.txt'
    source.write_text('Parties: Alpha and Beta. Effective date: 2026-09-21.', encoding='utf-8')
    parent_result, observed, chat_tools = contract_delivery_harness(
        source_path=source, output_slots=['document'],
        params={'output_slot_types': {'document': 'file'}},
    )
    task = observed['task']
    terminal = _terminal(observed['events'])
    assert task['output_slots'] == ['document']
    assert terminal['status'] == 'succeeded'
    assert parent_result['status'] == 'ok'
    assert chat_tools.get_subagent_status(task['title'])['task']['status'] == 'succeeded'
    artifacts = chat_tools.get_subagent_artifacts(task['title'], keys=['document'])['artifacts']
    assert len(artifacts) == 1
    assert parent_result['artifacts'] == artifacts
    assert artifacts[0]['content_type'] == 'file'
    delivered = artifacts[0]['value']
    assert delivered['size'] > 0
    assert [paragraph.text for paragraph in Document(delivered['path']).paragraphs][1] == source.read_text(
        encoding='utf-8',
    )
    print('BADCASE_EVIDENCE=' + json.dumps({
        'case': 'B', 'objective': task['objective'], 'output_slots': task['output_slots'],
        'contract': task['params']['output_slot_types'],
        'terminal': terminal['status'], 'failure_code': None,
        'artifact_count': len(artifacts), 'artifact_type': artifacts[0]['content_type'],
        'artifact_size': delivered['size'], 'document_readback': 'matched supplied input',
        'parent_status': parent_result['status'],
    }, ensure_ascii=False))


@pytest.fixture
def completion_case(monkeypatch, tmp_path):
    import lazymind.chat.engine.tools.subagent_chat_tools as chat_tools

    _install_fake_lazyllm(monkeypatch)
    _install_fake_translator(monkeypatch)
    monkeypatch.setattr(
        runner_mod, '_generate_display_plan',
        lambda *_args, **_kwargs: ['Prepare task', 'Execute task', 'Deliver result'],
    )
    monkeypatch.setattr(chat_tools, '_agentic_config', lambda: {'mode': 'manual'})
    created, contexts = [], []
    monkeypatch.setattr(chat_tools, '_write_agent_data', lambda tag, **kw: created.append(kw))
    monkeypatch.setattr(runner_mod, 'set_context', contexts.append)
    monkeypatch.setattr(runner_mod.subagent_tools, 'require_context', lambda: contexts[-1])
    model = MagicMock()
    monkeypatch.setattr(runner_mod, 'AutoModel', lambda **_: model)

    def run(*, params=None, output_slots=None, save=None, verdict=None,
            final='Delivered the requested result.', agent_type='document_generation',
            objective='Generate a PDF document and deliver the signed file.', events=(), stopped=None,
            previous_content_type=None, draft=None):
        model.share.return_value.return_value = json.dumps(verdict or {
            'completed': True, 'requires_artifact': True, 'artifact_keys': ['document'],
            'reason': 'The requested document was delivered.',
        })
        chat_tools.create_subagent(
            agent_type=agent_type, title='Task', objective=objective,
            params=params, output_slots=output_slots,
        )
        task = {**created[-1], 'id': created[-1]['task_id'], 'workspace_path': str(tmp_path)}
        if stopped:
            task['status'] = stopped
        if previous_content_type:
            previous = tmp_path / 'previous.pdf'
            previous.write_bytes(b'%PDF-1.4\nsynthetic prior deliverable\n')
            task['artifacts'] = [{'slot': 'document', 'content_type': previous_content_type,
                                  'seq': 1, 'value': {'path': str(previous), 'text': 'Prior result'}}]

        class Executor:
            async def stream(self, llm, plan):
                for event in events:
                    yield 'event', event
                if draft:
                    contexts[-1].write_draft('document', 'text', draft)
                if save:
                    key, content_type = save
                    value = 'Useful analysis.'
                    if content_type == 'file':
                        (tmp_path / 'document.pdf').write_bytes(b'%PDF-1.4\nsynthetic fixture\n')
                        value = 'document.pdf'
                    runner_mod.subagent_tools._save_artifact(key, value, content_type)
                yield 'final', final

        monkeypatch.setattr(runner_mod, 'AgentExecutor', Executor)
        raw = asyncio.run(_collect(runner_mod.run_subagent_stream(
            task['id'], task_spec=task, resume=bool(previous_content_type),
        )))
        result = _sse_to_events(raw)
        assert raw.endswith('data: [DONE]\n\n')
        return result, model, task

    return run


def _terminal(events):
    terminal = [event for event in events if event['type'] in {'done', 'error'}]
    assert len(terminal) == 1
    return terminal[0]


async def _run_completion_review_scenario(
    monkeypatch,
    tmp_path,
    model,
    *,
    agent_type='research',
    objective='Explain the analysis in plain text.',
    params=None,
    output_slots=None,
    artifacts=None,
    final='Analysis delivered.',
    before_final=None,
    on_chunk=None,
):
    fake_llm = MagicMock()
    fake_llm.globals._init_sid = lambda sid: None
    fake_llm.locals._init_sid = lambda sid: None
    monkeypatch.setattr(runner_mod, 'lazyllm', fake_llm)
    monkeypatch.setattr(runner_mod, 'AutoModel', lambda **_: model)
    monkeypatch.setattr(runner_mod, 'inject_model_config', lambda cfg: None)
    monkeypatch.setattr(runner_mod, 'inject_tool_config', lambda cfg: None)
    monkeypatch.setattr(runner_mod, 'inject_runtime_env', lambda cfg: None)
    monkeypatch.setattr(runner_mod, 'set_context', lambda ctx: None)
    monkeypatch.setattr(
        runner_mod,
        '_generate_display_plan',
        lambda *_args, **_kwargs: ['Prepare task', 'Execute task', 'Deliver result'],
    )
    _install_fake_translator(monkeypatch)

    class Executor:
        async def stream(self, _llm, _plan):
            if before_final is not None:
                await before_final()
            yield 'final', final

    monkeypatch.setattr(runner_mod, 'AgentExecutor', Executor)
    task = {
        'id': 'completion-review-task',
        'conversation_id': 'completion-review-conversation',
        'agent_type': agent_type,
        'objective': objective,
        'params': params or {},
        'workspace_path': str(tmp_path),
        'input_slots': [],
        'output_slots': output_slots or [],
        'artifacts': artifacts or [],
        'mode': 'auto',
    }
    raw = await _collect(
        runner_mod.run_subagent_stream(task['id'], task_spec=task, resume=bool(artifacts)),
        on_chunk=on_chunk,
    )
    return raw, _sse_to_events(raw)


@pytest.mark.asyncio
async def test_completion_review_does_not_block_event_loop(monkeypatch, tmp_path):
    verdict = json.dumps({
        'completed': True,
        'requires_artifact': False,
        'artifact_keys': [],
        'reason': 'Analysis delivered.',
    })

    evaluation_finished = threading.Event()

    class SlowModel:
        def share(self, **_kwargs):
            def call(_prompt):
                time.sleep(0.2)
                evaluation_finished.set()
                return verdict

            return call

    timer = None
    started = 0.0

    async def timer_elapsed():
        await asyncio.sleep(0.025)
        return asyncio.get_running_loop().time() - started, evaluation_finished.is_set()

    original_evaluate = runner_mod._evaluate_completion_async

    async def timed_evaluate(*args, **kwargs):
        nonlocal timer, started
        started = asyncio.get_running_loop().time()
        timer = asyncio.create_task(timer_elapsed())
        await asyncio.sleep(0)
        return await original_evaluate(*args, **kwargs)

    monkeypatch.setattr(runner_mod, '_evaluate_completion_async', timed_evaluate)
    _raw, events = await _run_completion_review_scenario(
        monkeypatch, tmp_path, SlowModel(),
    )

    assert timer is not None
    elapsed, evaluator_already_finished = await timer
    assert elapsed < 0.18
    assert evaluator_already_finished is False
    assert _terminal(events)['status'] == 'succeeded'


@pytest.mark.asyncio
@pytest.mark.parametrize('provider', ['qwen', 'openai'])
async def test_completion_review_supports_silent_streaming_models(monkeypatch, tmp_path, provider):
    calls = []

    class StreamingOnlyModel:
        def share(self, **kwargs):
            calls.append(kwargs)
            if kwargs.get('stream') is False:
                raise RuntimeError(f'{provider} provider requires streaming')
            sink = kwargs['stream']['_stream_sink']

            def call(_prompt):
                sink({'reasoning_content': 'PRIVATE_COMPLETION_REASONING'})
                return json.dumps({
                    'completed': True,
                    'requires_artifact': False,
                    'artifact_keys': [],
                    'reason': 'Analysis delivered.',
                })

            return call

    raw, events = await _run_completion_review_scenario(
        monkeypatch, tmp_path, StreamingOnlyModel(),
    )

    assert _terminal(events)['status'] == 'succeeded'
    assert calls and isinstance(calls[-1].get('stream'), dict)
    assert 'PRIVATE_COMPLETION_REASONING' not in raw


@pytest.mark.asyncio
async def test_missing_resumed_file_is_not_completion_evidence(monkeypatch, tmp_path):
    model = MagicMock()
    missing = tmp_path / 'deleted.pdf'
    raw, events = await _run_completion_review_scenario(
        monkeypatch,
        tmp_path,
        model,
        agent_type='document_generation',
        objective='Deliver the requested document.',
        params={'output_slot_types': {'document': 'file'}},
        output_slots=['document'],
        artifacts=[{
            'slot': 'document',
            'content_type': 'file',
            'seq': 1,
            'value': {'filename': 'deleted.pdf', 'path': str(missing), 'size': 12},
        }],
        final='The requested document was delivered.',
    )

    assert not missing.exists()
    assert _terminal(events)['status'] == 'failed', raw
    assert _terminal(events)['current_phase'] == 'missing_required_artifacts'
    model.share.assert_not_called()


@pytest.mark.asyncio
async def test_cancel_during_completion_review_discards_late_success(monkeypatch, tmp_path):
    evaluation_started = threading.Event()

    class SlowModel:
        def share(self, **_kwargs):
            def call(_prompt):
                evaluation_started.set()
                time.sleep(0.2)
                return json.dumps({
                    'completed': True,
                    'requires_artifact': False,
                    'artifact_keys': [],
                    'reason': 'Late success must be ignored.',
                })

            return call

    def cancel_check(_output):
        if evaluation_started.is_set():
            raise runner_mod.UserCancelledError('stopped during completion review')
        return False

    monkeypatch.setattr(runner_mod, 'make_cancel_stop_condition', lambda: cancel_check)
    evaluation_boundary = 0.0

    async def before_final():
        nonlocal evaluation_boundary
        evaluation_boundary = asyncio.get_running_loop().time()

    raw, events = await _run_completion_review_scenario(
        monkeypatch, tmp_path, SlowModel(), before_final=before_final,
    )

    assert asyncio.get_running_loop().time() - evaluation_boundary < 0.15
    assert _terminal(events)['status'] == 'interrupted'
    await asyncio.sleep(0.25)
    assert not any(event['status'] == 'succeeded' for event in events if 'status' in event)
    assert 'Late success must be ignored.' not in raw


@pytest.mark.asyncio
async def test_completion_review_waits_for_result_without_deadline(monkeypatch):
    release_evaluation = threading.Event()
    evaluation_started = threading.Event()

    class SlowModel:
        def share(self, **_kwargs):
            def call(_prompt):
                evaluation_started.set()
                assert release_evaluation.wait(timeout=2)
                return json.dumps({
                    'completed': True,
                    'requires_artifact': False,
                    'artifact_keys': [],
                    'reason': 'Analysis delivered after waiting.',
                })
            return call

    # Even a previously configured short deadline must no longer end the review.
    monkeypatch.setattr(runner_mod, '_cfg', {'subagent_completion_evaluation_timeout': 0.05})
    evaluation = asyncio.create_task(runner_mod._evaluate_completion_async(
        SlowModel(), 'Analyze the request', [], [], 'Analysis delivered.',
    ))
    try:
        assert await asyncio.to_thread(evaluation_started.wait, 1)
        await asyncio.sleep(0.15)
        assert not evaluation.done(), 'slow review must not be turned into task failure'
    finally:
        release_evaluation.set()
        result = await asyncio.wait_for(evaluation, timeout=1)
    assert result == (True, 'Analysis delivered after waiting.', '')


@pytest.mark.asyncio
async def test_valid_resumed_file_is_completion_evidence(monkeypatch, tmp_path):
    document = tmp_path / 'existing.pdf'
    document.write_bytes(b'%PDF-1.4\nexisting delivery\n')
    model = MagicMock()
    _raw, events = await _run_completion_review_scenario(
        monkeypatch,
        tmp_path,
        model,
        agent_type='document_generation',
        objective='Deliver the requested document.',
        params={'output_slot_types': {'document': 'file'}},
        output_slots=['document'],
        artifacts=[{
            'slot': 'document', 'content_type': 'file', 'seq': 1,
            'value': {'filename': document.name, 'path': str(document), 'size': document.stat().st_size},
        }],
        final='The requested document was delivered.',
    )

    assert _terminal(events)['status'] == 'succeeded'
    model.share.assert_not_called()


@pytest.mark.asyncio
@pytest.mark.parametrize('missing_index,expected', [(None, 'succeeded'), (1, 'failed')])
async def test_resumed_file_list_requires_every_file(monkeypatch, tmp_path, missing_index, expected):
    paths = [tmp_path / 'part-1.txt', tmp_path / 'part-2.txt']
    for index, path in enumerate(paths):
        if index != missing_index:
            path.write_text(f'part {index + 1}', encoding='utf-8')
    model = MagicMock()
    _raw, events = await _run_completion_review_scenario(
        monkeypatch,
        tmp_path,
        model,
        agent_type='document_generation',
        objective='Deliver every generated document.',
        params={'output_slot_types': {'documents': 'file'}},
        output_slots=['documents'],
        artifacts=[{
            'slot': 'documents', 'content_type': 'file_list', 'seq': 1,
            'value': {'paths': [str(path) for path in paths]},
        }],
        final='Every requested document was delivered.',
    )

    assert _terminal(events)['status'] == expected
    if expected == 'failed':
        assert _terminal(events)['current_phase'] == 'missing_required_artifacts'
    model.share.assert_not_called()


@pytest.mark.asyncio
@pytest.mark.parametrize('original_type', ['text', 'json'])
async def test_deleted_offloaded_content_is_not_completion_evidence(
    monkeypatch, tmp_path, original_type,
):
    missing = tmp_path / f'deleted-{original_type}.txt'
    model = MagicMock()
    _raw, events = await _run_completion_review_scenario(
        monkeypatch,
        tmp_path,
        model,
        objective='Deliver the requested report.',
        params={'output_slot_types': {'report': original_type}},
        output_slots=['report'],
        artifacts=[{
            'slot': 'report', 'content_type': 'file', 'seq': 1,
            'value': {'type': original_type, 'path': str(missing), 'size': 1024},
        }],
        final='The requested report was delivered.',
    )

    assert _terminal(events)['status'] == 'failed'
    assert _terminal(events)['current_phase'] == 'missing_required_artifacts'
    model.share.assert_not_called()


@pytest.mark.asyncio
async def test_artifact_v2_remote_blob_is_not_rejected_as_a_local_path(monkeypatch, tmp_path):
    model = MagicMock()
    _raw, events = await _run_completion_review_scenario(
        monkeypatch,
        tmp_path,
        model,
        agent_type='document_generation',
        objective='Deliver the requested document.',
        params={'output_slot_types': {'document': 'file'}},
        output_slots=['document'],
        artifacts=[{
            'slot': 'document', 'content_type': 'file', 'seq': 1,
            'v2_artifact_id': 'artifact-v2', 'v2_revision_id': 'revision-v2',
            'value': {'filename': 'report.pdf', 'url': 'https://files.example.test/report.pdf'},
        }],
        final='The requested document was delivered.',
    )

    assert _terminal(events)['status'] == 'succeeded'
    model.share.assert_not_called()


def _assert_silent_stream_call(model):
    assert model.share.call_count == 1
    stream = model.share.call_args.kwargs.get('stream')
    assert isinstance(stream, dict)
    assert callable(stream.get('_stream_sink'))


@pytest.mark.parametrize('params', [
    {'required_output_artifact_keys': ['document']},
    {'output_slot_types': {'document': 'file'}},
])
def test_structured_artifact_contract_without_slots_is_enforced(completion_case, params):
    events, model, task = completion_case(params=params)
    assert task['output_slots'] == ['document']
    assert _terminal(events)['status'] == 'failed'
    assert _terminal(events)['current_phase'] == 'missing_required_artifacts'
    model.share.assert_not_called()


@pytest.mark.parametrize('params', [{}, {'output_slot_types': {'document': 'file'}}])
def test_created_artifact_task_succeeds_with_saved_deliverable(completion_case, params):
    events, model, _ = completion_case(params=params, save=('document', 'file'))
    assert _terminal(events)['status'] == 'succeeded'
    artifacts = [event for event in events if event['type'] == 'artifact']
    assert artifacts and artifacts[0]['content_type'] == 'file'
    assert artifacts[0]['value']['size'] > 0
    assert model.share.call_count == (0 if params else 1)


@pytest.mark.parametrize('completed,expected', [(True, 'succeeded'), (False, 'failed')])
def test_uncontracted_text_task_is_evaluated_without_creating_files(completion_case, completed, expected):
    events, model, _ = completion_case(
        agent_type='research', objective='Explain the analysis in plain text.',
        final='The failure rate fell after retrying.' if completed else 'I cannot access the required input.',
        verdict={'completed': completed, 'requires_artifact': False, 'artifact_keys': [],
                 'reason': 'Analysis delivered.' if completed else 'Required input unavailable.'},
    )
    assert _terminal(events)['status'] == expected
    assert not any(event['type'] == 'artifact' for event in events)
    _assert_silent_stream_call(model)
    if not completed:
        assert _terminal(events)['current_phase'] == 'objective_incomplete'


def test_missing_artifact_overrides_positive_evaluator_verdict(completion_case):
    events, _, _ = completion_case()
    assert _terminal(events)['status'] == 'failed'
    assert _terminal(events)['current_phase'] == 'missing_required_artifacts'
    assert _terminal(events)['summary'] == 'Required deliverables were not saved as artifacts.'


def test_unrelated_artifact_does_not_complete_objective(completion_case):
    events, _, _ = completion_case(save=('notes', 'text'))
    assert _terminal(events)['status'] == 'failed'


def test_explicit_output_slots_cannot_be_waived_by_evaluator(completion_case):
    events, model, _ = completion_case(output_slots=['document'])
    assert _terminal(events)['status'] == 'failed'
    assert not any(event['type'] == 'artifact' for event in events)
    model.share.assert_not_called()


@pytest.mark.parametrize('completed,expected', [(True, 'succeeded'), (False, 'failed')])
def test_failure_words_trigger_review_not_an_automatic_failure(completion_case, completed, expected):
    events, model, _ = completion_case(
        output_slots=['document'], save=('document', 'file'),
        final='The first conversion failed.' if completed else 'I cannot complete the signature step.',
        verdict={'completed': completed, 'requires_artifact': True, 'artifact_keys': ['document'],
                 'reason': 'Retry succeeded.' if completed else 'Required signature was not completed.'},
    )
    assert _terminal(events)['status'] == expected
    _assert_silent_stream_call(model)


def test_recovered_tool_error_does_not_poison_completion(completion_case):
    events, model, _ = completion_case(
        output_slots=['document'], save=('document', 'file'), events=[
            {'tag': 'tool_results', 'tool_results': [{'id': '1', 'name': 'convert', 'result': {'status': 'failed'}}]},
            {'tag': 'tool_results', 'tool_results': [{'id': '2', 'name': 'convert', 'result': {'status': 'ok'}}]},
        ],
    )
    assert _terminal(events)['status'] == 'succeeded'
    model.share.assert_not_called()


@pytest.mark.parametrize('stopped', ['canceled', 'interrupted'])
def test_stopped_task_never_evaluates_or_succeeds(completion_case, stopped):
    events, model, _ = completion_case(stopped=stopped)
    assert _terminal(events)['status'] in {'canceled', 'interrupted'}
    assert not any(event['type'] == 'task_start' for event in events)
    model.share.assert_not_called()


@pytest.mark.parametrize('params,save,expected', [
    ({}, None, 'succeeded'),
    ({'required_output_artifact_keys': ['document'], 'legacy_tools': ['publish']}, None, 'failed'),
    ({'required_output_artifact_keys': ['document'], 'legacy_tools': ['publish']}, ('document', 'file'), 'succeeded'),
    ({'required_output_artifact_keys': ['document'], 'output_slot_types': {'document': 'text'}}, None, 'succeeded'),
    ({'required_output_artifact_keys': ['document'], 'output_slot_types': {'document': 'file'}}, None, 'failed'),
    ({'required_output_artifact_keys': ['document'],
      'workflow_runtime': {'publisher_owned_slots': ['document']}}, None, 'failed'),
])
def test_workflow_output_contracts_keep_their_existing_rules(completion_case, params, save, expected):
    events, model, _ = completion_case(
        agent_type='workflow_step', params=params, output_slots=['document'], save=save,
    )
    assert _terminal(events)['status'] == expected
    model.share.assert_not_called()


def test_text_cannot_replace_a_declared_file_artifact(completion_case):
    events, _, _ = completion_case(
        params={'output_slot_types': {'document': 'file'}}, save=('document', 'text'),
    )
    assert _terminal(events)['status'] == 'failed'
    assert not any(event['type'] == 'artifact' for event in events)


@pytest.mark.parametrize('content_type,expected', [('file', 'succeeded'), ('text', 'failed')])
def test_resume_rechecks_the_declared_artifact_type(completion_case, content_type, expected):
    events, model, _ = completion_case(
        params={'output_slot_types': {'document': 'file'}}, previous_content_type=content_type,
    )
    assert _terminal(events)['status'] == expected
    model.share.assert_not_called()


def test_required_text_draft_is_committed_before_completion(completion_case):
    events, model, _ = completion_case(output_slots=['document'], draft='The actual analysis.')
    assert _terminal(events)['status'] == 'succeeded'
    assert any(event['type'] == 'artifact' and event['value'].get('text') == 'The actual analysis.'
               for event in events)
    model.share.assert_not_called()


def test_workflow_final_incomplete_report_is_reviewed(completion_case):
    events, model, _ = completion_case(
        agent_type='workflow_step', final='I cannot complete the required action.',
        verdict={'completed': False, 'requires_artifact': False, 'artifact_keys': [],
                 'reason': 'The required action was not completed.'},
    )
    assert _terminal(events)['status'] == 'failed'
    assert _terminal(events)['current_phase'] == 'objective_incomplete'
    _assert_silent_stream_call(model)


def test_cancel_after_draft_publication_cannot_emit_success(monkeypatch, completion_case):
    calls = 0

    def check(_):
        nonlocal calls
        calls += 1
        if calls == 2:
            raise runner_mod.UserCancelledError('stopped after artifact publication')

    monkeypatch.setattr(runner_mod, 'make_cancel_stop_condition', lambda: check)
    events, model, _ = completion_case(output_slots=['document'], draft='Completed analysis.')
    assert any(event['type'] == 'artifact' for event in events)
    assert _terminal(events)['status'] == 'interrupted'
    model.share.assert_not_called()


@pytest.mark.parametrize('response', [
    'YES', '{}', '[]', '{"completed":"true"}',
    json.dumps({'completed': True, 'requires_artifact': False, 'artifact_keys': [], 'reason': ''}),
    json.dumps({'completed': True, 'requires_artifact': True, 'artifact_keys': 'document', 'reason': 'Done'}),
])
def test_invalid_completion_verdict_fails_closed(response):
    model = MagicMock()
    model.share.return_value.return_value = response
    completed, _, phase = runner_mod._evaluate_completion(model, 'objective', [], [], 'result')
    assert not completed
    assert phase == 'completion_evaluation_failed'


def test_completion_model_error_fails_closed():
    model = MagicMock()
    model.share.return_value.side_effect = RuntimeError('model unavailable')
    completed, _, phase = runner_mod._evaluate_completion(model, 'objective', [], [], 'result')
    assert not completed
    assert phase == 'completion_evaluation_failed'


@pytest.mark.parametrize('verdict,expected_phase', [
    ({'completed': True, 'requires_artifact': False, 'artifact_keys': [],
      'reason': 'Analysis delivered.'}, ''),
    ({'completed': False, 'requires_artifact': True, 'artifact_keys': [],
      'reason': 'The required PDF was not delivered.'}, 'missing_required_artifacts'),
])
def test_completion_review_keeps_reasoning_separate_from_json(monkeypatch, verdict, expected_phase):
    from lazyllm import OnlineChatModule

    llm = OnlineChatModule(source='deepseek', model='deepseek-v4-pro', api_key='test-key', stream=False)
    original_formatter = llm._formatter
    response = MagicMock()
    response.__enter__.return_value = response
    response.status_code = 200
    chunks = [
        {'role': 'assistant', 'content': ''},
        {'reasoning_content': 'PRIVATE_COMPLETION_REASONING'},
        {'content': json.dumps(verdict)},
    ]
    response.iter_lines.return_value = iter([
        ('data: ' + json.dumps({'choices': [{'index': 0, 'delta': delta}]})).encode()
        for delta in chunks
    ] + [
        b'data: {"choices":[{"index":0,"delta":{},"finish_reason":"stop"}]}',
        b'data: [DONE]',
    ])
    post = MagicMock(return_value=response)
    monkeypatch.setattr('requests.post', post)
    enqueue = MagicMock(side_effect=AssertionError('Completion review leaked into agent stream'))
    monkeypatch.setattr('lazyllm.FileSystemQueue.enqueue', enqueue)

    completed, summary, phase = runner_mod._evaluate_completion(llm, 'Task objective', [], [], 'Result')

    assert completed is verdict['completed']
    assert verdict['reason'] in summary
    assert phase == expected_phase
    assert post.call_args.kwargs['json']['stream'] is True
    assert llm._stream is False
    assert llm._formatter is original_formatter
    enqueue.assert_not_called()


@pytest.mark.parametrize('during_evaluation', [False, True])
def test_cancel_at_completion_boundary_cannot_be_revived(monkeypatch, completion_case, during_evaluation):
    canceled = not during_evaluation

    def check(_):
        if canceled:
            raise runner_mod.UserCancelledError('stopped by user')

    async def evaluate(**_):
        nonlocal canceled
        canceled = True
        return True, 'Late success verdict', ''

    monkeypatch.setattr(runner_mod, 'make_cancel_stop_condition', lambda: check)
    evaluate_mock = MagicMock(side_effect=evaluate)
    monkeypatch.setattr(runner_mod, '_evaluate_completion_async', evaluate_mock)
    events, _, _ = completion_case()
    assert _terminal(events)['status'] == 'interrupted'
    assert evaluate_mock.call_count == int(during_evaluation)


def test_subagent_closes_agent_before_disposing_database(monkeypatch):
    db = _install_fake_db(monkeypatch)
    _install_fake_lazyllm(monkeypatch)
    _install_fake_translator(monkeypatch)
    order = []
    db.dispose = lambda: order.append('dispose')
    monkeypatch.setattr(runner_mod, '_generate_display_plan', lambda *_: [])

    class Executor:
        async def stream(self, *_):
            try:
                yield 'event', {'tag': 'text', 'delta': 'ready ' * 100}
                await asyncio.Event().wait()
            finally:
                order.append('producer closed')

    monkeypatch.setattr(runner_mod, 'AgentExecutor', Executor)

    async def scenario():
        stream = runner_mod.run_subagent_stream(_DEFAULT_TASK_ID, task_spec=_DEFAULT_TASK)
        while 'ready ' not in await anext(stream):
            pass
        await stream.aclose()
        assert order == ['producer closed', 'dispose']

    asyncio.run(scenario())


def test_subagent_plan_preserves_extension_params_without_structured_duplicates(tmp_path):
    from lazymind.chat.engine.subagent.context import SubAgentContext

    ctx = SubAgentContext(
        task_id='task-params',
        conversation_id='conv-1',
        agent_type='test',
        objective='do something',
        params={
            'custom_plugin_option': 'keep-me',
            'count': 3,
            'history_files_per_turn': {'1': ['/tmp/input.txt']},
            'partial_indices': {'items': [0]},
            'required_output_artifact_keys': ['result'],
            'workflow_id': 'test-workflow',
            'session_id': 'session-1',
            'step_id': 'prompt',
            'user_input': 'run the whole workflow',
        },
        workspace_path=str(tmp_path),
        input_slots=[],
        output_slots=['result'],
        db=None,
        emit=lambda _event: None,
    )

    plan = runner_mod._build_subagent_plan(
        ctx,
        None,
        tools=[],
        tool_prompt_appendices={},
    )

    parameter_section = next(
        section for section in plan.prompt.sections
        if section.section_id == 'subagent_parameters'
    )
    assert 'custom_plugin_option: keep-me' in parameter_section.content
    assert 'count: 3' in parameter_section.content
    assert 'history_files_per_turn' not in parameter_section.content
    assert 'partial_indices' not in parameter_section.content
    assert 'required_output_artifact_keys' not in parameter_section.content
    assert 'workflow_id' not in parameter_section.content
    assert 'session_id' not in parameter_section.content
    assert 'user_input' not in parameter_section.content
    role_section = next(
        section for section in plan.prompt.sections
        if section.section_id == 'subagent_role'
    )
    assert 'must never ask the user a question' in role_section.content
    assert 'Never emit a fenced Markdown block with the language `editable`' in role_section.content
    assert all(section.section_id != 'editable_writing' for section in plan.prompt.sections)


def test_subagent_plan_uses_200_rounds_in_max_mode(tmp_path):
    import lazyllm
    from lazymind.chat.engine.subagent.context import SubAgentContext

    ctx = SubAgentContext(
        task_id='task-max', conversation_id='conv-1', agent_type='research',
        objective='deep research', params={'_thinking_depth': 'max'}, workspace_path=str(tmp_path),
        input_slots=[], output_slots=[], db=None, emit=lambda _event: None,
    )
    previous = lazyllm.globals.get('agentic_config')
    try:
        lazyllm.globals['agentic_config'] = {'thinking_depth': 'max'}
        with runner_mod._cfg.temp('agentic_expanded_max_rounds', 200):
            plan = runner_mod._build_subagent_plan(
                ctx, None, tools=[], tool_prompt_appendices={},
            )
    finally:
        lazyllm.globals['agentic_config'] = previous or {}

    assert plan.execution_options.max_retries == 199


def test_subagent_plan_forwards_llm_config_for_context_budget(tmp_path):
    from lazymind.chat.engine.subagent.context import SubAgentContext

    ctx = SubAgentContext(
        task_id='task-budget', conversation_id='conv-1', agent_type='workflow_step',
        objective='retrieve literature', params={}, workspace_path=str(tmp_path),
        input_slots=[], output_slots=[], db=None, emit=lambda _event: None,
    )
    llm_config = {'llm': {'source': 'deepseek', 'model': 'deepseek-v4-flash', 'max_input_tokens': '1M'}}
    plan = runner_mod._build_subagent_plan(
        ctx, None, tools=[], tool_prompt_appendices={}, llm_config=llm_config,
    )

    assert plan.execution_options.llm_config == llm_config


def test_workflow_step_uses_overflow_only_history_compactor(tmp_path):
    from lazymind.chat.engine.subagent.context import SubAgentContext

    ctx = SubAgentContext(
        task_id='task-workflow-budget', conversation_id='conv-1', agent_type='workflow_step',
        objective='retrieve literature', params={}, workspace_path=str(tmp_path),
        input_slots=[], output_slots=[], db=None, emit=lambda _event: None,
    )

    plan = runner_mod._build_subagent_plan(
        ctx, None, tools=[], tool_prompt_appendices={},
    )

    assert plan.execution_options.workspace == str(tmp_path)
    assert plan.execution_options.history_compactor is not None


def test_ordinary_subagent_keeps_default_history_compactor(tmp_path):
    from lazymind.chat.engine.subagent.context import SubAgentContext

    ctx = SubAgentContext(
        task_id='task-default-budget', conversation_id='conv-1', agent_type='research',
        objective='retrieve literature', params={}, workspace_path=str(tmp_path),
        input_slots=[], output_slots=[], db=None, emit=lambda _event: None,
    )

    plan = runner_mod._build_subagent_plan(
        ctx, None, tools=[], tool_prompt_appendices={},
    )

    assert plan.execution_options.history_compactor is None


def test_ordinary_subagent_enables_inherited_skill_runtime(tmp_path):
    from lazymind.chat.engine.subagent.context import SubAgentContext

    ctx = SubAgentContext(
        task_id='task-image-skill', conversation_id='conv-1', agent_type='image_generation',
        objective='generate a presentation background',
        params={'_inherited_skills': ['design/image-prompt-craft']},
        workspace_path=str(tmp_path), input_slots=[], output_slots=[], db=None,
        emit=lambda _event: None,
    )
    plan = runner_mod._build_subagent_plan(
        ctx, None, tools=[], tool_prompt_appendices={},
    )

    assert plan.execution_options.skills == ['design/image-prompt-craft']
    assert plan.execution_options.prompt_skills == []
    assert plan.execution_options.fs is runner_mod.FS
    assert plan.execution_options.skills_dir
    assert all(
        '_inherited_skills' not in section.content for section in plan.prompt.sections
    )


@pytest.mark.parametrize('retrieval', [False, True])
def test_workflow_step_keeps_skill_runtime_isolated(tmp_path, monkeypatch, retrieval):
    from lazymind.chat.engine.subagent.context import SubAgentContext

    monkeypatch.setitem(runner_mod.lazyllm.globals, 'agentic_config', {'enable_tool_retrieval': retrieval})
    ctx = SubAgentContext(
        task_id='task-workflow-skill', conversation_id='conv-1', agent_type='workflow_step',
        objective='generate a presentation background',
        params={'_inherited_skills': ['design/image-prompt-craft']},
        workspace_path=str(tmp_path), input_slots=[], output_slots=[], db=None,
        emit=lambda _event: None,
    )
    plan = runner_mod._build_subagent_plan(
        ctx, None, tools=[], tool_prompt_appendices={},
    )

    assert plan.execution_options.skills is None
    assert plan.execution_options.prompt_skills is None
    assert plan.execution_options.fs is None
    assert plan.execution_options.skills_dir is None
    assert plan.execution_options.preload_all_tools is True
    assert plan.execution_options.enable_builtin_tools is (False if retrieval else None)


# ---------------------------------------------------------------------------
# Test: task not found
# ---------------------------------------------------------------------------

def test_run_subagent_stream_task_not_found(monkeypatch):
    db = FakeDB(task=None)
    monkeypatch.setattr(
        runner_mod,
        'MemorySubAgentStore',
        lambda task_spec, initial_steps=None, artifacts=None: db,
    )

    async def run():
        return await _collect(runner_mod.run_subagent_stream(
            'bad-id', task_spec={**_DEFAULT_TASK},
        ))

    raw = asyncio.run(run())
    events = _sse_to_events(raw)
    assert events[0]['type'] == 'error'
    assert 'not found' in events[0]['message']
    assert raw.endswith('data: [DONE]\n\n')


# ---------------------------------------------------------------------------
# Test: happy path SSE sequence
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("display_plan", [["Read inputs", "Analyze data", "Write result"], None])
def test_run_subagent_stream_happy_path(monkeypatch, display_plan):
    db = _install_fake_db(monkeypatch)
    _install_fake_lazyllm(monkeypatch)
    _install_fake_build(monkeypatch)
    _install_fake_translator(monkeypatch)

    def generate_plan(*_args):
        if display_plan is None:
            raise ValueError('Model returned an invalid plan')
        return display_plan

    monkeypatch.setattr(runner_mod, '_generate_display_plan', generate_plan)

    # Simulate: text event → tool_calls → tool_results (triggers artifact emit) → text
    tool_calls_event = {'tag': 'tool_calls', 'tool_calls': [{'id': 'c1', 'name': 'save_artifacts', 'args': {}}]}
    tool_results_event = {'tag': 'tool_results', 'tool_results': [{'id': 'c1', 'name': 'save_artifacts', 'result': 'ok'}]}
    events = [
        {'tag': 'text', 'delta': 'Starting...'},
        tool_calls_event,
        tool_results_event,
        {'tag': 'text', 'delta': 'Done.'},
    ]
    _install_fake_drive(monkeypatch, events)

    # Patch ctx.saved_keys to return declared key so completeness check passes.
    real_ensure = runner_mod.SubAgentContext if hasattr(runner_mod, 'SubAgentContext') else None
    original_set_ctx = runner_mod.set_context

    ctx_holder: list = []

    def capturing_set_context(ctx):
        ctx_holder.append(ctx)
        # Pre-populate saved keys to pass completeness check.
        ctx.record_local_artifact('result', 'text', {'text': 'Completed result'}, 1)

    monkeypatch.setattr(runner_mod, 'set_context', capturing_set_context)

    async def run():
        return await _collect(runner_mod.run_subagent_stream(
            _DEFAULT_TASK_ID, task_spec={**_DEFAULT_TASK},
        ))

    raw = asyncio.run(run())
    events_out = _sse_to_events(raw)
    types_out = [e['type'] for e in events_out]

    assert types_out[0] == 'task_start'
    assert types_out[1] == 'progress'  # initial progress
    if display_plan:
        plan_events = [e for e in events_out if e['type'] == 'plan']
        if plan_events:  # A very short task may finish before the optional plan.
            assert plan_events[0]['steps'] == display_plan
            assert next(s for s in db.steps if s['role'] == 'plan')['content']['steps'] == display_plan
    else:
        assert 'plan' not in types_out

    assert 'text' in types_out
    assert 'progress' in types_out   # tool_results progress bump
    assert 'done' in types_out
    assert raw.endswith('data: [DONE]\n\n')


def test_subagent_runtime_env_refreshes_without_entering_plan_or_events(monkeypatch, tmp_path):
    import lazyllm

    task = {**_DEFAULT_TASK, 'workspace_path': str(tmp_path)}
    db = _install_fake_db(monkeypatch, task)
    monkeypatch.setattr(runner_mod, 'AutoModel', lambda model: 'fake_llm')
    monkeypatch.setattr(runner_mod, 'inject_model_config', lambda cfg: None)
    monkeypatch.setattr(runner_mod, 'inject_tool_config', lambda cfg: None)
    monkeypatch.setattr(
        runner_mod,
        'set_context',
        lambda ctx: ctx.record_local_artifact('result', 'text', {'text': 'Completed result'}, 1),
    )
    _install_fake_translator(monkeypatch)
    expected = {}

    class Executor:
        async def stream(self, llm, plan):
            assert lazyllm.globals._sid == _DEFAULT_TASK_ID
            assert lazyllm.globals['dynamic_env_vars'] == expected
            assert lazyllm.globals['conversation_env_overrides'] == {}
            assert 'synthetic-secret' not in repr(plan)
            yield 'final', 'task done'

    monkeypatch.setattr(runner_mod, 'AgentExecutor', Executor)

    async def run():
        nonlocal expected
        for resume, values in [(False, {'Mixed_API_KEY': 'synthetic-secret-one'}),
                               (True, {'Mixed_API_KEY': 'synthetic-secret-two'}), (True, None)]:
            lazyllm.globals._init_sid(sid=_DEFAULT_TASK_ID)
            lazyllm.globals['dynamic_env_vars'] = {'STALE_TOKEN': 'synthetic-secret-stale'}
            lazyllm.globals['conversation_env_overrides'] = {'Mixed_API_KEY': 'synthetic-secret-other-chat'}
            expected = values or {}
            raw = await _collect(runner_mod.run_subagent_stream(
                _DEFAULT_TASK_ID, task_spec=task, resume=resume, user_env_vars=values,
            ))
            assert any(event['type'] == 'done' for event in _sse_to_events(raw))
            assert 'synthetic-secret' not in raw
            assert 'synthetic-secret' not in repr(db.steps)

    asyncio.run(run())


@pytest.mark.asyncio
async def test_subagent_api_forwards_private_runtime_env(monkeypatch):
    from lazymind.chat.api.subagent_routes import run_subagent

    captured = {}

    async def stream(**kwargs):
        captured.update(kwargs)
        yield 'data: [DONE]\n\n'

    monkeypatch.setattr(runner_mod, 'run_subagent_stream', stream)
    response = await run_subagent(
        task_id=_DEFAULT_TASK_ID, task_spec=_DEFAULT_TASK,
        user_env_vars={'Mixed_API_KEY': 'synthetic-secret'},
    )
    await _collect(response.body_iterator)
    assert captured['user_env_vars'] == {'Mixed_API_KEY': 'synthetic-secret'}
    assert 'user_env_vars' not in captured['task_spec']


def test_tool_result_sends_separate_resume_safe_payload(monkeypatch):
    db = _install_fake_db(monkeypatch)
    _install_fake_lazyllm(monkeypatch)
    _install_fake_build(monkeypatch)
    _install_fake_translator(monkeypatch)
    full_result = 'x' * 3000
    _install_fake_drive(monkeypatch, [{
        'tag': 'tool_results',
        'tool_results': [{'id': 'c1', 'name': 'read_file', 'result': full_result}],
    }])

    def pre_save_ctx(ctx):
        ctx.record_local_artifact('result', 'text', {'text': 'Completed result'}, 1)

    monkeypatch.setattr(runner_mod, 'set_context', pre_save_ctx)

    async def run():
        return await _collect(runner_mod.run_subagent_stream(
            _DEFAULT_TASK_ID, task_spec={**_DEFAULT_TASK},
        ))

    events = _sse_to_events(asyncio.run(run()))
    result_event = next(event for event in events if event.get('type') == 'tool_results')
    assert result_event['tool_results'][0]['result'] == full_result[:2000]
    assert result_event['durable_tool_results'][0]['result'] == full_result
    assert db.steps[0]['content']['tool_results'][0]['result'] == full_result


# ---------------------------------------------------------------------------
# Test: missing artifact → error frame
# ---------------------------------------------------------------------------

def test_run_subagent_stream_missing_artifact_emits_error(monkeypatch):
    db = _install_fake_db(monkeypatch)
    _install_fake_lazyllm(monkeypatch)
    _install_fake_build(monkeypatch)
    _install_fake_translator(monkeypatch)
    _install_fake_drive(monkeypatch, [])
    # set_context does NOT pre-populate saved keys → completeness check fails

    async def run():
        return await _collect(runner_mod.run_subagent_stream(
            _DEFAULT_TASK_ID, task_spec={**_DEFAULT_TASK},
        ))

    raw = asyncio.run(run())
    events_out = _sse_to_events(raw)
    error_events = [e for e in events_out if e.get('type') == 'error']
    assert error_events, 'Expected an error event for missing artifact'
    assert 'result' in error_events[0]['message']
    assert raw.endswith('data: [DONE]\n\n')


# ---------------------------------------------------------------------------
# Test: AgentExecutor raises → outer except → error frame
# ---------------------------------------------------------------------------

def test_run_subagent_stream_agent_exception_emits_error(monkeypatch):
    _install_fake_db(monkeypatch)
    _install_fake_lazyllm(monkeypatch)
    _install_fake_build(monkeypatch)
    _install_fake_translator(monkeypatch)

    class ExplodingExecutor:
        async def stream(self, llm, plan):
            raise RuntimeError('llm exploded')
            yield  # pragma: no cover

    monkeypatch.setattr(runner_mod, 'AgentExecutor', ExplodingExecutor)

    async def run():
        return await _collect(runner_mod.run_subagent_stream(
            _DEFAULT_TASK_ID, task_spec={**_DEFAULT_TASK},
        ))

    raw = asyncio.run(run())
    events_out = _sse_to_events(raw)
    error_events = [e for e in events_out if e.get('type') == 'error']
    assert error_events
    assert 'llm exploded' in error_events[0]['message']


# ---------------------------------------------------------------------------
# Test: text/think frames from AgentEventFrameTranslator appear in SSE
# ---------------------------------------------------------------------------

def test_run_subagent_stream_text_think_events(monkeypatch):
    db = _install_fake_db(monkeypatch)
    _install_fake_lazyllm(monkeypatch)
    _install_fake_build(monkeypatch)
    _install_fake_translator(monkeypatch)

    events = [
        {'tag': 'think', 'delta': 'reasoning...'},
        {'tag': 'text', 'delta': 'answer'},
    ]
    _install_fake_drive(monkeypatch, events)

    # Pre-populate saved key.
    def pre_save_ctx(ctx):
        ctx.record_local_artifact('result', 'text', {'text': 'Completed result'}, 1)
    monkeypatch.setattr(runner_mod, 'set_context', pre_save_ctx)

    async def run():
        return await _collect(runner_mod.run_subagent_stream(
            _DEFAULT_TASK_ID, task_spec={**_DEFAULT_TASK},
        ))

    raw = asyncio.run(run())
    events_out = _sse_to_events(raw)
    types_out = [e['type'] for e in events_out]
    assert 'think' in types_out
    assert 'text' in types_out


def test_run_subagent_stream_coalesces_tiny_text_deltas(monkeypatch):
    from types import SimpleNamespace
    # Exercise byte coalescing independently from scheduler pauses under load.
    monkeypatch.setattr(runner_mod, 'time', SimpleNamespace(time=runner_mod.time.time, monotonic=lambda: 0.0))
    db = _install_fake_db(monkeypatch)
    _install_fake_lazyllm(monkeypatch)
    _install_fake_build(monkeypatch)
    _install_fake_translator(monkeypatch)
    _install_fake_drive(monkeypatch, [
        {'tag': 'text', 'delta': 'x'} for _ in range(1024)
    ])

    def pre_save_ctx(ctx):
        ctx.record_local_artifact('result', 'text', {'text': 'Completed result'}, 1)
    monkeypatch.setattr(runner_mod, 'set_context', pre_save_ctx)

    async def run():
        return await _collect(runner_mod.run_subagent_stream(
            _DEFAULT_TASK_ID, task_spec={**_DEFAULT_TASK},
        ))

    events_out = _sse_to_events(asyncio.run(run()))
    text_events = [event for event in events_out if event.get('type') == 'text']
    assert ''.join(event.get('text', '') for event in text_events) == 'x' * 1024
    assert len(text_events) <= 4
    assert len(db.steps) == 1


def test_workflow_tool_internal_text_is_not_forwarded(monkeypatch):
    workflow_task = {
        **_DEFAULT_TASK,
        'agent_type': 'workflow_step',
        'params': {'required_output_artifact_keys': ['result']},
    }
    _install_fake_db(monkeypatch, workflow_task)
    _install_fake_lazyllm(monkeypatch)
    _install_fake_build(monkeypatch)
    _install_fake_translator(monkeypatch)
    _install_fake_drive(monkeypatch, [
        {'tag': 'text', 'delta': 'Starting.'},
        {'tag': 'tool_calls', 'tool_calls': [
            {'id': 'ppt-1', 'name': 'ppt_generate_pages', 'args': {}},
        ]},
        {'tag': 'text', 'delta': '<html>large internal page output</html>'},
        {'tag': 'tool_results', 'tool_results': [
            {'id': 'ppt-1', 'name': 'ppt_generate_pages', 'result': 'ok'},
        ]},
        {'tag': 'text', 'delta': 'Finished.'},
    ])

    def pre_save_ctx(ctx):
        ctx.record_local_artifact('result', 'text', {'text': 'Completed result'}, 1)
    monkeypatch.setattr(runner_mod, 'set_context', pre_save_ctx)

    async def run():
        return await _collect(runner_mod.run_subagent_stream(
            _DEFAULT_TASK_ID, task_spec=workflow_task,
        ))

    events_out = _sse_to_events(asyncio.run(run()))
    visible_text = ''.join(
        event.get('text', '') for event in events_out if event.get('type') == 'text'
    )
    assert visible_text == 'Starting.Finished.'
    assert '<html>' not in visible_text


def test_workflow_tool_artifact_is_streamed_before_tool_returns(monkeypatch):
    workflow_task = {
        **_DEFAULT_TASK,
        'agent_type': 'workflow_step',
        'params': {'required_output_artifact_keys': ['result']},
    }
    _install_fake_db(monkeypatch, workflow_task)
    _install_fake_lazyllm(monkeypatch)
    _install_fake_build(monkeypatch)
    _install_fake_translator(monkeypatch)
    context_holder: Dict[str, Any] = {}

    def capture_context(ctx):
        context_holder['ctx'] = ctx
        ctx.record_local_artifact('result', 'text', {'text': 'Completed result'}, 1)

    monkeypatch.setattr(runner_mod, 'set_context', capture_context)

    class ProgressiveArtifactExecutor:
        async def stream(self, llm, plan):
            yield 'event', {
                'tag': 'tool_calls',
                'tool_calls': [{'id': 'ppt-1', 'name': 'ppt_generate_pages', 'args': {}}],
            }
            context_holder['ctx'].emit({
                'type': 'artifact',
                'slot': 'result',
                'content_type': 'text',
                'seq': 1,
                'value': {'text': '<html>page one</html>', 'list_index': 0},
            })
            # Model a publisher that keeps generating more pages after page one
            # has already been made available to the UI.
            await asyncio.sleep(0.01)
            yield 'event', {
                'tag': 'tool_results',
                'tool_results': [{'id': 'ppt-1', 'name': 'ppt_generate_pages', 'result': 'ok'}],
            }
            yield 'final', 'done'

    monkeypatch.setattr(runner_mod, 'AgentExecutor', ProgressiveArtifactExecutor)

    async def run():
        return await _collect(runner_mod.run_subagent_stream(
            _DEFAULT_TASK_ID, task_spec=workflow_task,
        ))

    events_out = _sse_to_events(asyncio.run(run()))
    event_types = [event.get('type') for event in events_out]

    assert event_types.count('artifact') == 1
    assert event_types.index('artifact') < event_types.index('tool_results')


def test_run_subagent_stream_emits_task_scoped_source_snapshot(monkeypatch):
    _install_fake_db(monkeypatch)
    _install_fake_lazyllm(monkeypatch)
    _install_fake_build(monkeypatch)
    _install_fake_translator(monkeypatch)
    _install_fake_drive(monkeypatch, [{
        'tag': 'tool_results',
        'tool_results': [{'id': 'search-1', 'name': 'web_search', 'result': 'ok'}],
    }])
    monkeypatch.setattr(
        runner_mod,
        'materialize_source_views',
        MagicMock(side_effect=[[], [{
            'index': '1.1',
            'source_type': 'external',
            'title': 'Example',
            'url': 'https://example.test',
            'source_roles': ['searched'],
        }], [{
            'index': '1.1',
            'source_type': 'external',
            'title': 'Example',
            'url': 'https://example.test',
            'source_roles': ['searched'],
        }]]),
    )

    def pre_save_ctx(ctx):
        ctx.record_local_artifact('result', 'text', {'text': 'Completed result'}, 1)

    monkeypatch.setattr(runner_mod, 'set_context', pre_save_ctx)

    async def run():
        return await _collect(runner_mod.run_subagent_stream(
            _DEFAULT_TASK_ID, task_spec={**_DEFAULT_TASK},
        ))

    events_out = _sse_to_events(asyncio.run(run()))
    source_events = [event for event in events_out if event.get('type') == 'sources']
    assert len(source_events) == 1
    assert source_events[0]['task_id'] == _DEFAULT_TASK_ID
    assert source_events[0]['sources'][0]['source_roles'] == ['searched']


# ---------------------------------------------------------------------------
# Test: _rebuild_history_from_steps pairing validation
# ---------------------------------------------------------------------------

def test_rebuild_history_valid_pairs():
    db = FakeDB()
    db.steps = [
        {'task_id': 't1', 'seq': 0, 'role': 'assistant',
         'content': {'text': '', 'tool_calls': [{'id': 'c1', 'name': 'tool_a', 'args': {}}]}},
        {'task_id': 't1', 'seq': 1, 'role': 'tool',
         'content': {'tool_results': [{
             'tool_call_id': 'c1',
             'name': 'tool_a',
             'result': {'path': '/tmp/result.txt', 'offset': 4},
         }]}},
    ]
    history = runner_mod._rebuild_history_from_steps(db, 't1')
    assert len(history) == 2
    assert history[0]['role'] == 'assistant'
    assert history[1]['role'] == 'tool'
    assert history[1]['tool_call_id'] == 'c1'
    assert history[1][runner_mod.TOOL_OBSERVATION_KEY] == {
        'version': 1,
        'ok': None,
        'value': {'path': '/tmp/result.txt', 'offset': 4},
        'error': '',
    }


def test_rebuild_history_skips_observation_for_string_tool_results():
    db = FakeDB()
    db.steps = [
        {'task_id': 't1', 'seq': 0, 'role': 'assistant',
         'content': {'text': '', 'tool_calls': [{'id': 'c1', 'name': 'tool_a', 'args': {}}]}},
        {'task_id': 't1', 'seq': 1, 'role': 'tool',
         'content': {'tool_results': [{
             'tool_call_id': 'c1',
             'name': 'tool_a',
             'result': 'plain tool output',
         }]}},
    ]
    history = runner_mod._rebuild_history_from_steps(db, 't1')
    assert history[1]['content'] == 'plain tool output'
    assert runner_mod.TOOL_OBSERVATION_KEY not in history[1]


def test_rebuild_history_orphan_tool_result_dropped():
    db = FakeDB()
    db.steps = [
        {'task_id': 't1', 'seq': 0, 'role': 'assistant',
         'content': {'text': '', 'tool_calls': [{'id': 'c1', 'name': 'tool_a', 'args': {}}]}},
        {'task_id': 't1', 'seq': 1, 'role': 'tool',
         'content': {'tool_results': [{'tool_call_id': 'WRONG_ID', 'name': 'tool_a', 'result': 'r'}]}},
    ]
    history = runner_mod._rebuild_history_from_steps(db, 't1')
    # Orphan tool result: assistant step should also be dropped (we stop at last complete boundary).
    assert all(h.get('role') != 'tool' for h in history)


def test_rebuild_history_no_steps_returns_empty():
    db = FakeDB()
    history = runner_mod._rebuild_history_from_steps(db, 't1')
    assert history == []


@pytest.mark.parametrize('resume,identity', [
    (False, {'task_id': _DEFAULT_TASK_ID, 'generation': 'launch-1'}),
    (True, {'task_id': _DEFAULT_TASK_ID, 'generation': 'launch-2'}),
    (True, None),
])
def test_fastapi_subagent_launch_identity_reaches_runner_privately(monkeypatch, tmp_path, resume, identity):
    import httpx
    from fastapi import FastAPI
    from lazymind.chat.api.subagent_routes import router

    task = {**_DEFAULT_TASK, 'workspace_path': str(tmp_path), 'output_slots': [], 'params': {
        'user_id': 'owner', '_workspace_execution': {'generation': 'later-persisted'},
        'parent_agentic_config': {'run_id': 'parent-run', '_workspace_execution': {'generation': 'parent'}},
    }}
    _install_fake_db(monkeypatch, task)
    _install_fake_lazyllm(monkeypatch)
    _install_fake_translator(monkeypatch)
    monkeypatch.setattr(runner_mod, '_resolve_runtime_tools', lambda *_: [])
    monkeypatch.setattr(runner_mod, '_build_subagent_tools', lambda *_, **__: [])
    configs, prompts = [], []
    runner_mod.lazyllm.globals.__setitem__.side_effect = lambda key, value: configs.append(value) if key == 'agentic_config' else None

    class Executor:
        async def stream(self, _llm, plan):
            prompts.append(plan.prompt.system_prompt + plan.prompt.current_input)
            yield 'final', 'done'

    monkeypatch.setattr(runner_mod, 'AgentExecutor', Executor)
    app = FastAPI()
    app.include_router(router)

    async def request():
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url='http://test') as client:
            return await client.post('/api/subagent/run', json={
                'task_id': _DEFAULT_TASK_ID,
                'task_spec': task,
                'initial_steps': [],
                'resume': resume,
                'workspace_execution': identity,
            })

    response = asyncio.run(request())
    assert response.status_code == 200
    assert any(item.get('status') == 'succeeded' for item in _sse_to_events(response.text))
    assert len(configs) == len(prompts) == 1
    assert configs[0]['_workspace_execution'] == (identity or {})
    for private in ['launch-1', 'launch-2', 'later-persisted', 'parent-run', 'body-params']:
        assert private not in prompts[0] and private not in response.text


def test_user_cancel_is_interrupted_instead_of_missing_output_failure(monkeypatch):
    _install_fake_db(monkeypatch)
    _install_fake_lazyllm(monkeypatch)
    _install_fake_build(monkeypatch)
    _install_fake_translator(monkeypatch)

    class Executor:
        async def stream(self, *_):
            raise runner_mod.UserCancelledError('stopped by user')
            yield  # Keep the same async iterator interface as AgentExecutor.

    monkeypatch.setattr(runner_mod, 'AgentExecutor', Executor)
    events = _sse_to_events(asyncio.run(_collect(runner_mod.run_subagent_stream(
        _DEFAULT_TASK_ID, task_spec={**_DEFAULT_TASK},
    ))))
    assert not any(event['type'] == 'error' for event in events)
    assert next(event for event in events if event['type'] == 'done')['status'] == 'interrupted'


def test_display_plan_is_model_generated_and_bounded():
    llm = MagicMock()
    llm.share.return_value.return_value = '```json\n["检索销售数据", "比较季度趋势", "整理分析报告"]\n```'
    assert runner_mod._generate_display_plan(llm, '分析季度销售') == [
        '检索销售数据', '比较季度趋势', '整理分析报告',
    ]
    assert '分析季度销售' in llm.share.return_value.call_args.args[0]
    llm.share.return_value.return_value = '["one", {"text": "two"}, "three"]'
    with pytest.raises(ValueError):
        runner_mod._generate_display_plan(llm, 'task')


@pytest.mark.parametrize('source', ['qwen', 'openai', 'deepseek'])
def test_display_plan_collects_stream_without_emitting_agent_events(monkeypatch, source):
    from lazyllm import OnlineChatModule

    llm = OnlineChatModule(source=source, model='qwq-plus', api_key='test-key', stream=False)
    response = MagicMock()
    response.__enter__.return_value = response
    response.status_code = 200
    chunks = [
        {'role': 'assistant', 'content': ''},
        {'reasoning_content': 'Private planning reasoning'},
        {'content': '["Read inputs",'},
        {'content': '"Analyze data",'},
        {'content': '"Write result"]'},
    ]
    response.iter_lines.return_value = iter([
        ('data: ' + json.dumps({'choices': [{'index': 0, 'delta': delta}]})).encode()
        for delta in chunks
    ] + [
        b'data: {"choices":[{"index":0,"delta":{},"finish_reason":"stop"}]}',
        b'data: [DONE]',
    ])
    post = MagicMock(return_value=response)
    monkeypatch.setattr('requests.post', post)
    enqueue = MagicMock(side_effect=AssertionError('Background plan leaked into agent stream'))
    monkeypatch.setattr('lazyllm.FileSystemQueue.enqueue', enqueue)

    assert runner_mod._generate_display_plan(llm, 'Analyze sales') == [
        'Read inputs', 'Analyze data', 'Write result',
    ]
    assert post.call_args.kwargs['json']['stream'] is True
    assert llm._stream is False
    enqueue.assert_not_called()


def test_display_plan_is_not_replayed_as_private_agent_history():
    db = FakeDB()
    db.append_step('t', 0, 'plan', {'steps': ['a', 'b', 'c']})
    assert runner_mod._rebuild_history_from_steps(db, 't') == []


def test_display_plan_preserves_current_step_contract_before_overall_request():
    llm = MagicMock()
    llm.share.return_value.return_value = '["识别受众与演示目标", "梳理页面要求与视觉约束", "整理需求简报和能力要求"]'
    scope = {
        'step_id': 'analyze_requirements',
        'prompt': 'Analyze the user requirements. Produce only the requirements brief and capability marker.',
        'acceptance_criteria': ['Audience and visual constraints are explicit'],
        'output_slots': ['requirement_analysis', 'ppt_capability_requirements'],
    }
    steps = runner_mod._generate_display_plan(llm, '制作三页 PPT，收集图片、生成底图、建立大纲。', scope)
    prompt = llm.share.return_value.call_args.args[0]
    payload = json.loads(prompt.split('\n', 1)[1])
    assert payload['current_step'] == 'analyze_requirements'
    assert payload['current_step_contract'] == scope['prompt']
    assert payload['current_step_outputs'] == scope['output_slots']
    assert payload['current_step_acceptance'] == scope['acceptance_criteria']
    assert '制作三页' in payload['task_context_only']
    assert 'CURRENT SUBTASK ONLY' in prompt
    assert len(steps) == 3


def test_unconfigured_retrieval_is_not_exposed_by_subagent(monkeypatch):
    import lazyllm
    original = lazyllm.globals.get('agentic_config')
    lazyllm.globals['agentic_config'] = {}
    try:
        old_auth = lazyllm.globals.config['dynamic_tool_auth']
        lazyllm.globals.config['dynamic_tool_auth'] = {}
        try:
            monkeypatch.setattr(runner_mod, 'load_workflow_tools', lambda *_: {})
            assert runner_mod._resolve_runtime_tools(['kb', 'web_search']) == []
            assert runner_mod.subagent_tools.list_knowledge_bases not in runner_mod._build_subagent_tools([])
        finally:
            lazyllm.globals.config['dynamic_tool_auth'] = old_auth
    finally:
        lazyllm.globals['agentic_config'] = original or {}


def test_display_plan_does_not_delay_execution_and_is_persisted_live(monkeypatch):
    import threading
    task = {**_DEFAULT_TASK, 'params': {'required_output_artifact_keys': []}, 'output_artifact_keys': []}
    db = _install_fake_db(monkeypatch, task)
    _install_fake_lazyllm(monkeypatch)
    _install_fake_build(monkeypatch)
    _install_fake_translator(monkeypatch)
    started = threading.Event()
    outline = ['Read inputs', 'Analyze data', 'Write result']

    def generate(*_):
        assert started.wait(1), 'execution was blocked by display-plan generation'
        return outline

    class Executor:
        async def stream(self, *_):
            started.set()
            for _ in range(100):
                if any(s['role'] == 'plan' for s in db.steps):
                    break
                await asyncio.sleep(0.005)
            yield 'final', 'done'

    monkeypatch.setattr(runner_mod, '_generate_display_plan', generate)
    monkeypatch.setattr(runner_mod, 'AgentExecutor', Executor)
    async def complete(*_, **__):
        return True, 'done', ''

    monkeypatch.setattr(runner_mod, '_evaluate_completion_async', complete)
    raw = asyncio.run(_collect(runner_mod.run_subagent_stream(_DEFAULT_TASK_ID, task_spec=task)))
    events = _sse_to_events(raw)
    assert next(e for e in events if e['type'] == 'plan')['steps'] == outline
    assert any(e['type'] == 'done' for e in events)
    assert len({s['seq'] for s in db.steps}) == len(db.steps)
