import importlib.util
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile

from pptx import Presentation
from pptx.util import Pt

spec = importlib.util.spec_from_file_location('pptx_template_fill', (Path(__file__).parents[2] / 'algorithm')
                                              / 'lazymind/chat/engine/tools/file_resources/pptx_template_fill.py')
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)


class TemplateFillTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.source = Path(self.temp.name) / 'template.pptx'
        self.output = Path(self.temp.name) / 'output.pptx'
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = 'Sample title'
        run = slide.placeholders[1].text_frame.paragraphs[0].add_run()
        run.text = 'Sample achievement'
        run.font.bold = True
        run.font.size = Pt(20)
        prs.slides.add_slide(prs.slide_layouts[1]).shapes.title.text = 'Keep page'
        prs.save(self.source)
        self.pages = [module.inspect_template(self.source, p) for p in (1, 2)]
        self.edits = [{'id': s['id'], 'expected_text': s['text'], 'action': 'keep'}
                      for p in self.pages for s in p['slots']]

    def fill(self):
        return module.fill_template(self.source, self.output, self.pages[0]['sha256'], self.edits)

    def test_preserves_package_parts_and_styles(self):
        original = self.source.read_bytes()
        self.edits[0].update(action='replace', text='My <work> & results')
        result = self.fill()
        self.assertEqual(result['slide_count'], 2)
        self.assertEqual(self.source.read_bytes(), original)
        with ZipFile(self.source) as before, ZipFile(self.output) as after:
            self.assertEqual(before.namelist(), after.namelist())
            for name in before.namelist():
                if name != 'ppt/slides/slide1.xml':
                    self.assertEqual(before.read(name), after.read(name), name)
        actual = Presentation(self.output)
        self.assertEqual(actual.slides[0].shapes.title.text, 'My <work> & results')
        run = actual.slides[0].placeholders[1].text_frame.paragraphs[0].runs[0]
        self.assertTrue(run.font.bold)
        self.assertEqual(run.font.size, Pt(20))

    def test_missing_or_duplicate_slots_rejected(self):
        self.edits.pop()
        with self.assertRaises(ValueError):
            self.fill()
        self.assertFalse(self.output.exists())
        self.edits.append(self.edits[0])
        with self.assertRaises(ValueError):
            self.fill()

    def test_empty_placeholder_can_be_filled(self):
        empty = next(s for s in self.pages[1]['slots'] if s['empty_paragraph'])
        edit = next(e for e in self.edits if e['id'] == empty['id'])
        edit.update(action='replace', text='New work summary')
        self.fill()
        actual = Presentation(self.output)
        self.assertIn('New work summary', '\n'.join(s.text for s in actual.slides[1].shapes if s.has_text_frame))

    def test_stale_text_and_hash_rejected(self):
        self.edits[0]['expected_text'] = 'wrong'
        with self.assertRaises(ValueError):
            self.fill()
        with self.assertRaises(ValueError):
            module.fill_template(self.source, self.output, 'wrong', self.edits)

    def test_source_overwrite_rejected(self):
        with self.assertRaises(ValueError):
            module.fill_template(self.source, self.source, self.pages[0]['sha256'], self.edits)

    def test_overflow_warning_and_page_bounds(self):
        self.edits[0].update(action='replace', text='Long content ' * 30)
        with self.assertRaisesRegex(ValueError, 'does not fit'):
            self.fill()
        self.assertFalse(self.output.exists())
        with self.assertRaises(ValueError):
            module.inspect_template(self.source, 0)

    def test_inherited_geometry_and_quality_contract(self):
        self.assertIn('quality_contract', self.pages[0])
        self.assertTrue(any(r['status'] != 'unknown' and r['width_pt'] > 0
                            for r in self.pages[0]['layout_checks']))

    def test_blank_lines_count_towards_capacity(self):
        self.edits[0].update(action='replace', text='Short' + '\n' * 30)
        with self.assertRaisesRegex(ValueError, 'does not fit'):
            self.fill()
        self.assertFalse(self.output.exists())

    def test_whole_paragraph_newlines_create_paragraphs_with_style(self):
        target = next(e for e in self.edits if e['expected_text'] == 'Sample achievement')
        target.update(action='replace', text='First line\nSecond line')
        self.fill()
        with ZipFile(self.output) as archive:
            root = module._xml(archive.read('ppt/slides/slide1.xml'))
            self.assertEqual(len(root.findall('.//a:br', module.NS)), 0)
            self.assertTrue(all('\n' not in (n.text or '') for n in root.findall('.//a:t', module.NS)))
        paragraphs = Presentation(self.output).slides[0].placeholders[1].text_frame.paragraphs
        self.assertEqual([p.text for p in paragraphs], ['First line', 'Second line'])
        self.assertTrue(all(r.font.bold for p in paragraphs for r in p.runs))
        self.assertTrue(all(r.font.size == Pt(20) for p in paragraphs for r in p.runs))

    def test_empty_placeholder_materializes_inherited_size_and_theme_font(self):
        empty = next(s for s in self.pages[1]['slots'] if s['empty_paragraph'])
        next(e for e in self.edits if e['id'] == empty['id']).update(action='replace', text='中文正文\n第二段')
        stats = self.fill()
        with ZipFile(self.output) as archive:
            root = module._xml(archive.read('ppt/slides/slide2.xml'))
            for run in root.xpath('.//a:r[a:t="中文正文" or a:t="第二段"]', namespaces=module.NS):
                rp = run.find('a:rPr', module.NS)
                self.assertIsNotNone(rp.get('sz'))
                self.assertFalse(rp.find('a:latin', module.NS).get('typeface').startswith('+'))
        self.assertFalse(stats['post_generation_tools_required'])
        self.assertEqual(stats['generation_status'], 'completed')

    def test_inline_rich_text_keeps_inline_breaks(self):
        prs = Presentation(self.source)
        p = prs.slides[0].placeholders[1].text_frame.paragraphs[0]
        p.add_run().text = ' suffix'
        prs.save(self.source)
        pages = [module.inspect_template(self.source, i) for i in (1, 2)]
        edits = [{'id': s['id'], 'expected_text': s['text'], 'action': 'replace' if s['text']
                  == 'Sample achievement' else 'keep', 'text': 'First\nSecond'} for page in pages for s in page['slots']]
        module.fill_template(self.source, self.output, pages[0]['sha256'], edits)
        p = Presentation(self.output).slides[0].placeholders[1].text_frame.paragraphs
        self.assertEqual(len(p), 1)
        self.assertIn('suffix', p[0].text)

    def test_auto_sized_table_has_no_negative_capacity(self):
        from pptx.util import Inches
        prs = Presentation(self.source)
        table = prs.slides[0].shapes.add_table(1, 1, Inches(1), Inches(2), Inches(2), Inches(.4)).table
        table.cell(0, 0).text = 'Label'
        table._tbl.find('a:tr', module.NS).set('h', '0')
        prs.save(self.source)
        report = next(r for r in module.inspect_template(self.source)['layout_checks'] if r['shape_id'] == 'table:1:1:1')
        self.assertEqual(report['status'], 'unknown')
        self.assertNotIn('height_pt', report)

    def test_unavailable_fonts_are_disclosed_without_substitution(self):
        from unittest.mock import patch
        from lazymind.chat.engine.tools.file_resources.pptx_text_style import font_availability
        with patch(
            'lazymind.chat.engine.tools.file_resources.pptx_text_style._installed_fonts', return_value=frozenset()
        ), \
                patch('lazymind.chat.engine.tools.file_resources.font_library.local_fonts', return_value=[]):
            result = font_availability({'Missing Template Font'})
        self.assertEqual(result['unavailable_locally'], ['Missing Template Font'])
        self.assertFalse(result['substitution_applied'])

    def test_table_cell_capacity_rejected_before_writing(self):
        from pptx.util import Inches
        prs = Presentation(self.source)
        table = prs.slides[0].shapes.add_table(1, 1, Inches(1), Inches(2), Inches(2), Inches(.4)).table
        table.cell(0, 0).text = 'Label'
        prs.save(self.source)
        pages = [module.inspect_template(self.source, p) for p in (1, 2)]
        edits = [{'id': s['id'], 'expected_text': s['text'], 'action': 'replace' if s['text'] == 'Label' else 'keep',
                  'text': 'Too much text ' * 100} for p in pages for s in p['slots']]
        with self.assertRaisesRegex(ValueError, 'table:1:1:1'):
            module.fill_template(self.source, self.output, pages[0]['sha256'], edits)
        self.assertFalse(self.output.exists())

    def test_shared_cache_reuses_all_pages_and_binds_one_template(self):
        from unittest.mock import patch
        workspace = Path(self.temp.name) / 'conversation'
        first = module.inspect_template(self.source, 1, cache_workspace=workspace)
        self.assertFalse(first['cache_hit'])
        with patch.object(module, '_load', side_effect=AssertionError('must not parse again')):
            second = module.inspect_template(self.source, 2, cache_workspace=workspace)
        self.assertTrue(second['cache_hit'])
        self.assertEqual(second['cached_pages'], 2)
        other = Path(self.temp.name) / 'other.pptx'
        prs = Presentation(self.source)
        prs.slides[0].shapes.title.text = 'Other template'
        prs.save(other)
        with self.assertRaisesRegex(ValueError, 'different PPTX template'):
            module.inspect_template(other, 1, cache_workspace=workspace)
        isolated = module.inspect_template(other, 1, cache_workspace=Path(self.temp.name) / 'another-conversation')
        self.assertFalse(isolated['cache_hit'])

    def test_single_line_does_not_charge_trailing_line_spacing(self):
        from pptx.util import Inches
        prs = Presentation(self.source)
        box = prs.slides[0].shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(.5))
        p = box.text_frame.paragraphs[0]
        p.text = 'Short'
        p.runs[0].font.size = Pt(24)
        p.line_spacing = 1.5
        prs.save(self.source)
        pages = [module.inspect_template(self.source, i) for i in (1, 2)]
        edits = [{'id': s['id'], 'expected_text': s['text'], 'action': 'replace' if s['text'] == 'Short' else 'keep',
                  'text': 'New title'} for page in pages for s in page['slots']]
        module.fill_template(self.source, self.output, pages[0]['sha256'], edits)
        self.assertTrue(self.output.exists())

    def test_cache_corruption_reparsed_and_invalid_page_does_not_bind(self):
        workspace = Path(self.temp.name) / 'conversation'
        with self.assertRaises(ValueError):
            module.inspect_template(self.source, 99, cache_workspace=workspace)
        self.assertFalse((workspace / 'pptx-analysis/template.json').exists())
        first = module.inspect_template(self.source, 1, cache_workspace=workspace)
        cache = workspace / 'pptx-analysis' / (first['sha256'] + '.json')
        cache.write_text('broken')
        self.assertFalse(module.inspect_template(self.source, 1, cache_workspace=workspace)['cache_hit'])


if __name__ == '__main__':
    unittest.main()
