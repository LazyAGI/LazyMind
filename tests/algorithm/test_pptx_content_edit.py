from zipfile import ZipFile
import hashlib
from types import SimpleNamespace
from unittest.mock import Mock

import pytest
from pptx import Presentation
from lazymind.chat.engine.tools.file_resources.pptx_content_edit import inspect_content, edit_content


@pytest.fixture
def deck(tmp_path):
    path = tmp_path / 'source.pptx'
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[1]).shapes.title.text = '2026/11/30'
    second = prs.slides.add_slide(prs.slide_layouts[1])
    second.shapes.title.text = 'Results'
    second.notes_slide.notes_text_frame.text = 'Existing notes'
    prs.save(path)
    return path


def test_notes_and_date_preserve_other_slides_and_source(deck, tmp_path):
    original = deck.read_bytes()
    info = inspect_content(deck)
    slot = next(s for s in info['pages'][0]['slots'] if s['text'] == '2026/11/30')
    output = tmp_path / 'edited.pptx'
    notes = [{'page': p, 'text': f'第{p}页详细讲稿\n背景、具体工作与下一页过渡。' * 30} for p in (1, 2)]
    stats = edit_content(deck, output, info['sha256'],
                         [{'id': slot['id'], 'expected_text': slot['text'], 'text': '2025/11/30'}], notes, True)
    assert stats['updated_text_pages'] == [1]
    assert stats['updated_notes_pages'] == [1, 2]
    assert deck.read_bytes() == original
    result = Presentation(output)
    assert result.slides[0].shapes.title.text == '2025/11/30'
    assert [s.notes_slide.notes_text_frame.text for s in result.slides] == [n['text'] for n in notes]
    with ZipFile(deck) as before, ZipFile(output) as after:
        for name in before.namelist():
            if name.startswith(('ppt/slideMasters/', 'ppt/slideLayouts/', 'ppt/theme/',
                               'ppt/media/')) or name == 'ppt/slides/slide2.xml':
                assert before.read(name) == after.read(name), name


def test_page_text_match_accepts_frozen_tool_arguments(deck, tmp_path):
    from types import MappingProxyType
    out = tmp_path / 'frozen.pptx'
    edit_content(deck, out, inspect_content(deck)['sha256'],
                 (MappingProxyType({'page': 1, 'expected_text': '2026/11/30', 'text': '2025/11/30'}),),
                 (MappingProxyType({'page': 1, 'text': '详细讲稿'}),))
    assert Presentation(out).slides[0].shapes.title.text == '2025/11/30'
    assert Presentation(out).slides[0].notes_slide.notes_text_frame.text == '详细讲稿'


def test_ambiguous_page_text_never_guesses(deck, tmp_path):
    prs = Presentation(deck)
    from pptx.util import Inches
    prs.slides[0].shapes.add_textbox(Inches(1), Inches(3), Inches(3), Inches(1)).text = '2026/11/30'
    prs.save(deck)
    out = tmp_path / 'ambiguous.pptx'
    with pytest.raises(ValueError, match='matches 2 slots'):
        edit_content(deck, out, inspect_content(deck)['sha256'],
                     [{'page': 1, 'expected_text': '2026/11/30', 'text': '2025/11/30'}])
    assert not out.exists()


def test_notes_only_keeps_every_slide_xml_and_repeated_edit(deck, tmp_path):
    info = inspect_content(deck)
    first, second = tmp_path / 'first.pptx', tmp_path / 'second.pptx'
    edit_content(deck, first, info['sha256'], speaker_notes=[{'page': 2, 'text': 'Updated'}])
    edit_content(first, second, inspect_content(first)['sha256'], speaker_notes=[{'page': 2, 'text': 'Again'}])
    with ZipFile(deck) as before, ZipFile(second) as after:
        for name in before.namelist():
            if not name.startswith('ppt/notesSlides/notesSlide'):
                assert before.read(name) == after.read(name), name
    assert inspect_content(second)['pages'][1]['speaker_notes'] == 'Again'


@pytest.mark.parametrize('notes,all_pages', [([{'page': 3, 'text': 'x'}], False),
                                             ([{'page': 1, 'text': 'x'}], True), ([
                                                 {'page': 1, 'text': 'x'}, {'page': 1, 'text': 'y'}], False),
                                             ([{'page': 1, 'text': ''}, {'page': 2, 'text': 'x'}], True)])
def test_bad_notes_rejected_without_output(deck, tmp_path, notes, all_pages):
    output = tmp_path / 'bad.pptx'
    with pytest.raises(ValueError):
        edit_content(deck, output, inspect_content(deck)['sha256'], speaker_notes=notes, require_all_notes=all_pages)
    assert not output.exists()


def test_stale_source_and_text_rejected(deck, tmp_path):
    output = tmp_path / 'bad.pptx'
    with pytest.raises(ValueError, match='Deck changed'):
        edit_content(deck, output, 'stale', speaker_notes=[{'page': 1, 'text': 'x'}])
    info = inspect_content(deck)
    with pytest.raises(ValueError, match='Stale text'):
        edit_content(deck, output, info['sha256'], [{'id': '1:1', 'expected_text': 'wrong', 'text': 'new'}])
    assert not output.exists()


@pytest.mark.parametrize('id_key', ['id', 'artifact_id'])
def test_generated_deck_tool_publishes_selected_revision_with_cas(deck, tmp_path, monkeypatch, id_key):
    from lazymind.chat.engine.subagent import tools
    from lazymind.chat.engine.tools.file_resources.attachment_edit import AttachmentEditDraft
    target = {id_key: 'revision-5', 'revision': 5, 'draft_version': 2, 'slot': 'native_pptx',
              'selected': True, 'validity': 'effective', 'value': {'path': str(deck)}}
    monkeypatch.setattr(tools, '_public_workflow_artifacts', lambda *_: [target])
    monkeypatch.setattr(AttachmentEditDraft, 'for_current_conversation', lambda _: SimpleNamespace(
        root=str(tmp_path), effective_path=str(deck), draft_path=str(tmp_path / 'draft.pptx')))
    client = Mock()
    client.patch_artifact.return_value.result = {
        'id': 'revision-6', 'revision': 6,
        'value': {'storage': 'inline_base64', 'content_base64': 'NEVER_EXPOSE_FILE_BYTES'},
    }
    monkeypatch.setattr(tools, '_workflow_client', lambda: client)
    info = tools.edit_presentation('native_pptx', 'inspect_content', workflow_session_id='session')
    result = tools.edit_presentation('native_pptx', 'edit_content', workflow_session_id='session',
                                     source_sha256=info['sha256'], speaker_notes=[{'page': 2, 'text': 'Detailed notes'}])
    args, kwargs = client.patch_artifact.call_args
    assert args[:2] == ('revision-5', 5)
    assert kwargs['base_draft_version'] == 2
    assert args[2]['storage'] == 'inline_base64'
    assert result['workflow_artifact']['revision'] == 6
    assert 'value' not in result['workflow_artifact']
    assert 'NEVER_EXPOSE_FILE_BYTES' not in str(result)
    assert hashlib.sha256(deck.read_bytes()).hexdigest() == info['sha256']


def test_sdk_passes_draft_version(monkeypatch):
    from lazymind.workflow_sdk import WorkflowClient
    transport = Mock()
    client = WorkflowClient('http://localhost', 'owner', transport=transport)
    monkeypatch.setattr(client, '_decode', lambda r: r)
    client.patch_artifact('r1', 3, {'text': 'x'}, base_draft_version=2)
    assert transport.patch.call_args.kwargs['json']['base_draft_version'] == 2
