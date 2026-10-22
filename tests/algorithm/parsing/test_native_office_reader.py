import pytest
from lazyllm.module.module import ModuleExecutionError
from lazyllm.tools.rag.doc_node import DocNode
from lazymind.parsing.engine.readers.native_office import NativeOfficeReader
from lazymind.parsing.engine.transform import GeneralParser


@pytest.mark.parametrize('suffix', ['.docx', '.pptx'])
@pytest.mark.parametrize('ocr_type', ['none', 'paddleocr', 'mineru'])
def test_native_office_produces_chunkable_nodes(tmp_path, suffix, ocr_type):
    from docx import Document
    from pptx import Presentation
    from pptx.util import Inches
    from lazyllm.tools.rag.readers.reader_config_inject import inject_reader_config

    path = tmp_path / ('document' + suffix)
    text = 'native knowledge base content\n' * 100
    if suffix == '.docx':
        document = Document()
        document.add_paragraph(text)
    else:
        document = Presentation()
        slide = document.slides.add_slide(document.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = text
    document.save(path)
    try:
        inject_reader_config(ocr_config={'ocr_type': ocr_type, 'ocr_url': 'http://self-hosted-ocr'})
        nodes = NativeOfficeReader()(path, extra_info={'source': 'upload'}, use_cache=False)
        assert nodes and all(isinstance(node, DocNode) for node in nodes)
        assert all(node.metadata['source'] == 'upload' for node in nodes)
        assert 'native knowledge base content' in nodes[0].text
        chunks = GeneralParser(max_length=128, split_by='\n').forward(nodes[0])
        assert len(chunks) > 1
        assert all('native knowledge base content' in chunk.text for chunk in chunks)
    finally:
        inject_reader_config(ocr_config={})


def test_official_mineru_presentation_route_is_preserved():
    calls = []
    expected = [DocNode(text='OCR presentation')]

    def ocr(file, **kwargs):
        calls.append(file)
        return expected

    reader = NativeOfficeReader(presentation_reader=ocr)
    assert reader('deck.pptx', extra_info={'ocr_type': 'mineru'}, use_cache=False) == expected
    assert calls == ['deck.pptx']


def test_native_office_rejects_legacy_files():
    with pytest.raises(ModuleExecutionError, match='convert legacy files'):
        NativeOfficeReader()('legacy.doc', use_cache=False)
