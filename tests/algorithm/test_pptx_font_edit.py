from zipfile import ZipFile
from unittest.mock import Mock
from types import SimpleNamespace
import pytest
from pptx import Presentation
from pptx.util import Inches, Pt
from lazymind.chat.engine.tools.file_resources.pptx_content_edit import inspect_content
from lazymind.chat.engine.tools.file_resources.pptx_font_edit import edit_fonts


@pytest.fixture
def deck(tmp_path, monkeypatch):
    from lazymind.chat.engine.tools.file_resources import presentation_edit
    monkeypatch.setattr(presentation_edit, 'available_font_families', lambda: {'Test Font'})
    prs = Presentation()
    for p in range(2):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f'Title {p}'
        box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(1))
        run = box.text_frame.paragraphs[0].add_run()
        run.text = '正文 text'
        run.font.size, run.font.bold = Pt(20), True
        slide.notes_slide.notes_text_frame.text = 'Keep notes'
    path = tmp_path / 'source.pptx'
    prs.save(path)
    return path


@pytest.mark.parametrize('scope', ['all', 'page', 'slot', 'title'])
def test_font_scope_preserves_unselected_parts(deck, tmp_path, scope):
    info = inspect_content(deck)
    slot = next(s['id'] for s in info['pages'][0]['slots'] if s['text'] == '正文 text')
    kwargs = {'pages': [1]} if scope in {'page', 'title'} else {}
    if scope == 'slot':
        kwargs['slot_ids'] = [slot]
    if scope == 'title':
        kwargs['text_scope'] = 'titles'
    out = tmp_path / 'edited.pptx'
    stats = edit_fonts(deck, out, info['sha256'], 'Test Font', **kwargs)
    assert stats['updated_font_pages'] == ([1, 2] if scope == 'all' else [1])
    result = Presentation(out)
    for before, after in zip(Presentation(deck).slides, result.slides):
        assert [s.text for s in before.shapes if s.has_text_frame] == [s.text for s in after.shapes if s.has_text_frame]
        run = after.shapes[-1].text_frame.paragraphs[0].runs[0]
        assert run.font.size == Pt(20) and run.font.bold is True
    with ZipFile(deck) as before, ZipFile(out) as after:
        for name in before.namelist():
            if name not in ({'ppt/slides/slide1.xml', 'ppt/slides/slide2.xml'}
                            if scope == 'all' else {'ppt/slides/slide1.xml'}):
                assert before.read(name) == after.read(name), name
    assert result.slides[0].shapes[-1].text_frame.paragraphs[0].runs[0].font.name == (
        None if scope == 'title' else 'Test Font')


@pytest.mark.parametrize('kwargs', [{'pages': [99]}, {'pages': [1.5]},
                         {'slot_ids': ['1:999']}, {'slot_ids': []}, {'text_scope': 'guess'}])
def test_invalid_scope_never_publishes(deck, tmp_path, kwargs):
    out = tmp_path / 'invalid.pptx'
    with pytest.raises(ValueError):
        edit_fonts(deck, out, inspect_content(deck)['sha256'], 'Test Font', **kwargs)
    assert not out.exists()


def test_missing_font_and_stale_hash_fail(deck, tmp_path):
    out = tmp_path / 'invalid.pptx'
    with pytest.raises(ValueError, match='不可用'):
        edit_fonts(deck, out, inspect_content(deck)['sha256'], 'Not Installed')
    with pytest.raises(ValueError, match='Deck changed'):
        edit_fonts(deck, out, 'stale', 'Test Font')
    assert not out.exists()


@pytest.mark.parametrize('formatting', [False, True])
def test_native_font_edit_publishes_compact_receipt(deck, tmp_path, monkeypatch, formatting):
    from lazymind.chat.engine.subagent import tools
    monkeypatch.setattr(tools, '_public_workflow_artifacts', lambda *a: [{
        'artifact_id': 'a1', 'revision': 3, 'draft_version': 2, 'selected': True,
        'validity': 'effective', 'value': {'path': str(deck)}}])
    monkeypatch.setattr(
        tools.AttachmentEditDraft,
        'for_current_conversation',
        lambda p: SimpleNamespace(
            root=str(tmp_path),
            effective_path=p))
    client = Mock()
    client.patch_artifact.return_value.result = {
        'artifact_id': 'a2', 'revision': 4, 'value': {
            'content_base64': 'HIDDEN'}}
    monkeypatch.setattr(tools, '_workflow_client', lambda: client)
    args = {'style_edits': [{'pages': [1], 'bold': False, 'italic': True, 'font_size_pt': 18}]
            } if formatting else {'target_font': 'Test Font', 'pages': [1]}
    result = tools.edit_presentation('native_pptx', 'format_text' if formatting else 'replace_fonts', **args,
                                     source_sha256=inspect_content(deck)['sha256'], workflow_session_id='session')
    assert result['updated_style_pages' if formatting else 'updated_font_pages'] == [1]
    assert result['workflow_artifact']['revision'] == 4
    assert 'HIDDEN' not in str(result)
    assert client.patch_artifact.call_args.kwargs['base_draft_version'] == 2
    if formatting:
        import base64
        from io import BytesIO
        payload = client.patch_artifact.call_args.args[2]['content_base64']
        generated = Presentation(BytesIO(base64.b64decode(payload)))
        run = generated.slides[0].shapes[-1].text_frame.paragraphs[0].runs[0]
        assert run.font.bold is False and run.font.italic is True and run.font.size == Pt(18)


def test_font_download_is_opt_in_and_allowlisted(monkeypatch):
    from lazymind.chat.engine.tools.file_resources import font_library, presentation_edit
    monkeypatch.setattr(presentation_edit, 'available_font_families', lambda: {'Test Font'})
    assert font_library.ensure_font('test font', True) == 'Test Font'
    for name, download in [('Noto Serif SC', False), ('宋体', True), ('https://localhost/font.ttf', True)]:
        with pytest.raises(ValueError, match='不可用'):
            font_library.ensure_font(name, download)


def test_native_format_properties_preserve_content_and_other_parts(deck, tmp_path, monkeypatch):
    from lazymind.chat.engine.tools.file_resources import pptx_font_edit
    info = inspect_content(deck)
    slot = next(s['id'] for s in info['pages'][0]['slots'] if s['text'] == '正文 text')
    monkeypatch.setattr(pptx_font_edit, 'ensure_font', lambda *a: pytest.fail('format-only must not resolve fonts'))
    out = tmp_path / 'formatted.pptx'
    pptx_font_edit.edit_style_batch(deck, out, info['sha256'], [{
        'slot_ids': [slot], 'font_size_pt': 18, 'bold': False,
        'italic': True, 'underline': True, 'text_color': '#123ABC'}])
    run = Presentation(out).slides[0].shapes[-1].text_frame.paragraphs[0].runs[0]
    assert run.text == '正文 text'
    assert run.font.size == Pt(18) and run.font.bold is False
    assert run.font.italic is True and run.font.underline is True
    assert str(run.font.color.rgb) == '123ABC'
    assert run.font.name is None
    with ZipFile(deck) as before, ZipFile(out) as after:
        for name in before.namelist():
            if name != 'ppt/slides/slide1.xml':
                assert before.read(name) == after.read(name), name


def test_style_batch_handles_multiple_scopes_and_explicit_false(deck, tmp_path):
    from lazymind.chat.engine.tools.file_resources.pptx_font_edit import edit_style_batch
    from types import MappingProxyType
    info = inspect_content(deck)
    slot = next(s['id'] for s in info['pages'][1]['slots'] if s['text'] == '正文 text')
    out = tmp_path / 'batch.pptx'
    stats = edit_style_batch(deck, out, info['sha256'], (
        MappingProxyType({'pages': (1,), 'text_scope': 'titles', 'font_size_pt': 24, 'bold': True}),
        MappingProxyType({'slot_ids': (slot,), 'italic': True, 'bold': False, 'underline': False}),
    ))
    result = Presentation(out)
    assert stats['updated_style_pages'] == [1, 2]
    title = result.slides[0].shapes.title.text_frame.paragraphs[0].runs[0]
    assert title.font.size == Pt(24) and title.font.bold is True
    body = result.slides[1].shapes[-1].text_frame.paragraphs[0].runs[0]
    assert body.font.italic is True and body.font.bold is False and body.font.underline is False
    assert body.font.size == Pt(20)


def test_style_batch_failure_is_atomic_and_size_overflow_is_blocked(deck, tmp_path):
    from lazymind.chat.engine.tools.file_resources.pptx_font_edit import edit_style_batch
    out = tmp_path / 'failed.pptx'
    digest = inspect_content(deck)['sha256']
    with pytest.raises(ValueError, match='not found'):
        edit_style_batch(deck, out, digest, [{'pages': [1], 'italic': True}, {'slot_ids': ['2:999'], 'bold': False}])
    assert not out.exists()
    with pytest.raises(ValueError, match='overflow'):
        edit_style_batch(deck, out, digest, [{'pages': [1], 'font_size_pt': 400}])
    assert not out.exists()
    assert inspect_content(deck)['sha256'] == digest


@pytest.mark.parametrize('values', [{'bold': 'false'}, {'italic': 1}, {'underline': 'none'},
                                    {'font_size_pt': True}, {'font_size_pt': float('nan')}, {'font_size_pt': 0},
                                    {'font_size_pt': 401}, {'text_color': 'red'}, {'text_color': '#12345Z'}])
def test_text_format_validation_rejects_invalid_values(values):
    from lazymind.chat.engine.tools.file_resources.text_format import text_format
    with pytest.raises(ValueError):
        text_format(**values)


def test_download_validates_and_registers_font_with_license(tmp_path, monkeypatch):
    from contextlib import contextmanager
    from fontTools.fontBuilder import FontBuilder
    from fontTools.pens.ttGlyphPen import TTGlyphPen
    from lazymind.chat.engine.tools.file_resources import font_library, presentation_edit
    import httpx
    builder = FontBuilder(1000, isTTF=True)
    builder.setupGlyphOrder(['.notdef'])
    builder.setupCharacterMap({})
    builder.setupGlyf({'.notdef': TTGlyphPen(None).glyph()})
    builder.setupHorizontalMetrics({'.notdef': (500, 0)})
    builder.setupHorizontalHeader(ascent=800, descent=-200)
    builder.setupNameTable({'familyName': 'Download Test', 'styleName': 'Regular',
                           'fullName': 'Download Test Regular', 'psName': 'DownloadTest-Regular'})
    builder.setupOS2()
    builder.setupPost()
    source = tmp_path / 'test.ttf'
    builder.save(source)
    payload = source.read_bytes()
    monkeypatch.setenv('LAZYMIND_FONT_DIR', str(tmp_path / 'library'))
    monkeypatch.setattr(font_library, 'DOWNLOAD_FONTS', {'Download Test': 'downloadtest/font.ttf'})
    monkeypatch.setattr(presentation_edit, 'available_font_families',
                        lambda: {f['family'] for f in font_library.local_fonts()})
    response = SimpleNamespace(raise_for_status=lambda: None, iter_bytes=lambda: iter([payload]))

    @contextmanager
    def stream(*args):
        yield response
    client = Mock()
    client.__enter__ = Mock(return_value=client)
    client.__exit__ = Mock(return_value=False)
    client.get.return_value = SimpleNamespace(raise_for_status=lambda: None, text='SIL OPEN FONT LICENSE Version 1.1')
    client.stream = stream
    factory = Mock(return_value=client)
    monkeypatch.setattr(httpx, 'Client', factory)
    assert font_library.ensure_font('Download Test', True) == 'Download Test'
    assert len(font_library.local_fonts()) == 1
    assert list((tmp_path / 'library').glob('*.OFL.txt'))
    assert font_library.ensure_font('Download Test', True) == 'Download Test'
    factory.assert_called_once_with(timeout=60, follow_redirects=False)
