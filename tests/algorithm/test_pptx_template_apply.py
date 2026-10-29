from pathlib import Path
import tempfile
from zipfile import ZipFile, ZIP_DEFLATED

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches

from lazymind.chat.engine.tools.file_resources.pptx_template_apply import apply_pptx_template


A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
NS = {'a': A}


def _patch_template_theme(path: Path):
    with ZipFile(path) as archive:
        infos = archive.infolist()
        data = {item.filename: archive.read(item) for item in infos}
    root = etree.fromstring(data['ppt/theme/theme1.xml'])
    root.find('.//a:majorFont/a:latin', NS).set('typeface', 'Template Major')
    root.find('.//a:majorFont/a:ea', NS).set('typeface', 'Template Major')
    root.find('.//a:minorFont/a:latin', NS).set('typeface', 'Template Body')
    root.find('.//a:minorFont/a:ea', NS).set('typeface', 'Template Body')
    accent = root.find('.//a:clrScheme/a:accent1/*', NS)
    accent.set('val', 'AA5500')
    data['ppt/theme/theme1.xml'] = etree.tostring(root, xml_declaration=True, encoding='UTF-8')
    with ZipFile(path, 'w', compression=ZIP_DEFLATED) as archive:
        originals = {item.filename: item for item in infos}
        for name, raw in data.items():
            archive.writestr(originals.get(name, name), raw)


def _text(deck: Path):
    prs = Presentation(deck)
    return [[shape.text for shape in slide.shapes if shape.has_text_frame] for slide in prs.slides]


def test_apply_template_preserves_content_and_source(tmp_path):
    source, template, output = (tmp_path / name for name in ('source.pptx', 'template.pptx', 'output.pptx'))
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = 'Research title'
    slide.placeholders[1].text = 'Method and evidence'
    for shape in (slide.shapes.title, slide.placeholders[1]):
        shape.text_frame.paragraphs[0].runs[0].font.name = 'Source Font'
    rectangle = slide.shapes.add_shape(1, Inches(8), Inches(5), Inches(1), Inches(1))
    rectangle.fill.solid()
    rectangle.fill.fore_color.rgb = RGBColor(0x44, 0x72, 0xC4)
    prs.slides.add_slide(prs.slide_layouts[1]).shapes.title.text = 'Conclusion'
    prs.save(source)

    target = Presentation()
    sample = target.slides.add_slide(target.slide_layouts[1])
    sample.shapes.title.text = 'Template title'
    sample.shapes.title.left = Inches(1.3)
    sample.background.fill.solid()
    sample.background.fill.fore_color.rgb = RGBColor(0xFA, 0xF2, 0xE8)
    target.save(template)
    _patch_template_theme(template)

    original = source.read_bytes()
    before_text = _text(source)
    result = apply_pptx_template(source, template, output)

    assert source.read_bytes() == original
    assert _text(output) == before_text
    assert result['slide_count'] == 2
    assert result['theme_parts_replaced'] >= 1
    assert result['font_replacements'] > 0
    assert result['original_unchanged'] and result['template_unchanged']
    updated = Presentation(output)
    assert updated.slides[0].shapes.title.text_frame.paragraphs[0].runs[0].font.name == 'Template Major'
    assert updated.slides[0].placeholders[1].text_frame.paragraphs[0].runs[0].font.name == 'Template Body'


def test_apply_template_rejects_non_pptx_and_same_file(tmp_path):
    source = tmp_path / 'source.pptx'
    Presentation().save(source)
    with tempfile.NamedTemporaryFile(suffix='.pdf') as other:
        try:
            apply_pptx_template(source, Path(other.name), tmp_path / 'out.pptx')
        except ValueError as exc:
            assert 'template must be a .pptx' in str(exc)
        else:
            raise AssertionError('expected non-PPTX template rejection')
    try:
        apply_pptx_template(source, source, tmp_path / 'out.pptx')
    except ValueError as exc:
        assert 'must be different' in str(exc)
    else:
        raise AssertionError('expected identical source/template rejection')
